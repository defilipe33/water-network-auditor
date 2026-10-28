"""Entrega 1.8: percursos completos, índice operacional e relatórios sintéticos."""
import copy,json,sqlite3,tempfile,unittest,shutil
from pathlib import Path
from unittest.mock import patch
from test_auditoria_fronteiras import APP
from test_corredores_recuperacao import pacote
from test_recuperacao_catalogo import base_sintetica
from auditoria_topologica_watergems.corredores import calcular
from auditoria_topologica_watergems import exportacao_watergems as ex,gestao_catalogo as gc,relatorio_investigacao as rel
from auditoria_topologica_watergems.gerenciamento import Catalogo
from auditoria_topologica_watergems.localizacao_relatorio import Logradouros,Geocodificador,sanear_fonte
from qgis.core import QgsGeometry,QgsVectorLayer,QgsFeature,QgsProject


def exemplo():
    p=pacote([(101,1,2),(102,2,3),(103,2,4)])
    p['resultado']=dict(meta=dict(uid='ig-sintetica',topologia=dict(execucao_uid='topo-1',projeto='modelo',cenario=1),parametros=dict(crs='EPSG:31983',espacial_habilitado=True)))
    r=calcular(p,'ESTRUTURAL',[1],[3],comprimentos={101:10,102:20,103:5})
    h=dict(obj_uid='conj-1',topo_uid='topo-1',modo='ESTRUTURAL',vinculos=r['elementos'],dados=dict(criterios=r))
    return p,h


class ExportacaoCompletaTest(unittest.TestCase):
    def test_junctions_percurso_sem_ramo(self):
        p,h=exemplo();r=ex.preparar(p,h);self.assertEqual({e['DomainElementID'] for e in r['elementos']},{1,2,3,101,102});self.assertEqual([e['Elemento'] for e in r['percursos']],[101,102])
    def test_acumulado(self):
        p,h=exemplo();r=ex.preparar(p,h);self.assertEqual([e['ComprimentoAcumulado_m'] for e in r['percursos']],[10,30])
    def test_exclusao_manual_nao_inventa_continuidade(self):
        p,h=exemplo();h['vinculos']=h['vinculos'][1:];r=ex.preparar(p,h);self.assertNotIn(101,{e['DomainElementID'] for e in r['elementos']});self.assertEqual(r['percursos'][0]['Ordem'],1);self.assertTrue(r['avisos'])
    def test_rotas_compartilhadas_uniao_distinta(self):
        p,h=exemplo();r=calcular(p,'ESTRUTURAL',[1],[3,4],comprimentos={101:10,102:20,103:5});h.update(vinculos=r['elementos'],dados=dict(criterios=r));r=ex.preparar(p,h);self.assertEqual(len(r['elementos']),7);self.assertEqual(sum(e['Elemento']==101 for e in r['percursos']),2)
    def test_versao_incorreta_bloqueia(self):
        p,h=exemplo();h['topo_uid']='outra'
        with self.assertRaises(ValueError):ex.preparar(p,h)
    def test_uid_incompativel(self):
        p,h=exemplo();h['vinculos'][0]['dados']['uid']='outro'
        with self.assertRaises(ValueError):ex.preparar(p,h)
    def test_elemento_ausente(self):
        p,h=exemplo();del p['topologia']['modelo']['elementos'][101]
        with self.assertRaises(ValueError):ex.preparar(p,h)
    def test_duplicado_bloqueia(self):
        p,h=exemplo();h['vinculos']+=h['vinculos'][:1]
        with self.assertRaises(ValueError):ex.preparar(p,h)
    def test_legacy_sem_ordem(self):
        p,h=exemplo();h['dados']['criterios']='legado';r=ex.preparar(p,h);self.assertFalse(r['percursos']);self.assertTrue(r['avisos'])
    def test_id_nao_validado(self):
        p,h=exemplo()
        with self.assertRaises(ValueError):ex.consultas(ex.preparar(p,h),p['topologia']['modelo'],'ID','DomainElementID')
    def test_id_validado_por_tipo(self):
        p,h=exemplo();q=ex.consultas(ex.preparar(p,h),p['topologia']['modelo'],'ID','ElementID',True);self.assertEqual(len(q),2);self.assertTrue(all(q['expressao'].startswith('ElementID IN (') for q in q))
    def test_label_aspas_zero_textual(self):
        p,h=exemplo();p['topologia']['modelo']['elementos'][101]['label']="001'O";q=ex.consultas(ex.preparar(p,h),p['topologia']['modelo']);self.assertIn("'001''O'",str(q))
    def test_label_repetido(self):
        p,h=exemplo();p['topologia']['modelo']['elementos'][103]['label']=p['topologia']['modelo']['elementos'][101]['label']
        with self.assertRaises(ValueError):ex.consultas(ex.preparar(p,h),p['topologia']['modelo'])
    def test_lotes_preservam_todos(self):
        p,h=exemplo();d=ex.preparar(p,h);qs=ex.consultas(d,p['topologia']['modelo'],max_itens=1);self.assertEqual(sum(q['quantidade'] for q in qs),len(d['elementos']));self.assertEqual(len(qs),5)
    def test_limite_sem_truncar(self):
        p,h=exemplo();p['topologia']['modelo']['elementos'][101]['label']='x'*100
        with self.assertRaises(ValueError):ex.consultas(ex.preparar(p,h),p['topologia']['modelo'],max_caracteres=40)
    def test_conferencia(self):
        p,h=exemplo();d=ex.preparar(p,h);r=ex.conferir(d,d['elementos'][1:]+[d['elementos'][1]],'topo-1');self.assertEqual(len(r['ausentes']),1);self.assertEqual(len(r['duplicados']),1)
    def test_conferencia_versao(self):
        p,h=exemplo()
        with self.assertRaises(ValueError):ex.conferir(ex.preparar(p,h),[],'outra')
    def test_exportacao_atomica(self):
        p,h=exemplo()
        with tempfile.TemporaryDirectory() as t:
            out=Path(t)/'nova';ex.exportar(out,p,h);self.assertTrue((out/'PERCURSOS.csv').exists())
            with self.assertRaises(FileExistsError):ex.exportar(out,p,h)
    def test_falha_sem_parcial(self):
        p,h=exemplo();p['topologia']['modelo']['elementos'][101]['label']=None
        with tempfile.TemporaryDirectory() as t:
            out=Path(t)/'nova'
            with self.assertRaises(ValueError):ex.exportar(out,p,h)
            self.assertFalse(out.exists());self.assertFalse(list(Path(t).glob('.exportacao_*')))


class CatalogoGestaoTest(unittest.TestCase):
    def setUp(self):self.tmp=tempfile.TemporaryDirectory();self.p=Path(self.tmp.name);self.cat=Catalogo(self.p);self.b=base_sintetica(self.p/'base.gpkg');self.cat.registrar(self.b);self.r=self.cat.listar()[0]
    def tearDown(self):self.tmp.cleanup()
    def test_nome_nao_altera_fonte(self):
        b=self.b.read_bytes();gc.alterar(self.cat,self.r,'RENOMEAR','Outro nome');self.assertEqual(self.cat.listar()[0]['nome'],'Outro nome');self.assertEqual(b,self.b.read_bytes())
    def test_arquivo_fisico_preservado_ao_arquivar(self):gc.alterar(self.cat,self.r,'ARQUIVAR');self.assertTrue(self.b.exists());self.assertEqual(self.cat.listar()[0]['estado_catalogo'],'ARQUIVADA')
    def test_restaurar(self):gc.alterar(self.cat,self.r,'ARQUIVAR');gc.alterar(self.cat,self.r,'RESTAURAR');self.assertEqual(self.cat.listar()[0]['estado_catalogo'],'ATIVA')
    def test_retirar_e_reexaminar_nao_reinclui(self):gc.alterar(self.cat,self.r,'RETIRAR');self.cat.varrer();self.assertEqual(self.cat.listar(),[]);self.assertTrue(self.b.exists())
    def test_adicionar_restaura(self):gc.alterar(self.cat,self.r,'RETIRAR');gc.adicionar(self.cat,self.b);self.assertEqual(len(self.cat.listar()),1)
    def test_nome_igual_uuid_diferente(self):b=base_sintetica(self.p/'outra.gpkg',uid='outro');self.cat.registrar(b);self.assertEqual(len(self.cat.listar()),2)
    def test_copia_identica(self):shutil.copyfile(self.b,self.p/'copia.gpkg');self.cat.varrer();r=self.cat.listar();self.assertEqual(len(r),1);self.assertEqual(r[0]['classe_arquivo'],'COPIA_IDENTICA')
    def test_copia_divergente(self):base_sintetica(self.p/'divergente.gpkg',historico='Outro');self.cat.varrer();self.assertTrue(all(r['conflito_identidade'] for r in self.cat.listar()))
    def test_backup_separado(self):base_sintetica(self.p/'base_antes_conjuntos_ab.gpkg',historico='antes');self.cat.varrer();self.assertTrue(any(r['classe_arquivo']=='BACKUP_TECNICO' for r in self.cat.listar()))
    def test_indisponivel(self):self.b.rename(self.p/'movida.gpkg');self.assertEqual(self.cat.listar()[0]['classe_arquivo'],'INDISPONIVEL')
    def test_reabertura_overlay(self):gc.alterar(self.cat,self.r,'RENOMEAR','Novo');self.assertEqual(Catalogo(self.p).listar()[0]['nome'],'Novo')
    def test_fisico_sem_confirmacao(self):
        with self.assertRaises(ValueError):gc.recuperar_arquivo(self.cat,self.r)
    def test_base_topologica_compartilhada_nao_move(self):
        with self.assertRaises(ValueError):gc.recuperar_arquivo(self.cat,self.r,confirmado=True)
        self.assertTrue(self.b.exists())
    def test_contexto_ativo_bloqueado(self):self.assertTrue(gc.dependencias(self.cat,self.r,[str(self.b)])['bloqueios'])


class LocalizacaoMapaTest(unittest.TestCase):
    def ruas(self):
        c=QgsVectorLayer('LineString?crs=EPSG:31983&field=nome:string','Ruas sintéticas','memory')
        for nome,wkt in [('Rua A','LINESTRING(0 0,100 0)'),('Rua B','LINESTRING(50 -50,50 50)')]:
            f=QgsFeature(c.fields());f['nome']=nome;f.setGeometry(QgsGeometry.fromWkt(wkt));c.dataProvider().addFeatures([f])
        return c
    def test_rua_proxima_nao_endereco(self):r=Logradouros(self.ruas(),'nome','EPSG:31983').consultar(QgsGeometry.fromWkt('POINT(40 2)'));self.assertIn('não é endereço exato',r['texto']);self.assertEqual(r['candidatos'][0]['nome'],'Rua A')
    def test_cruzamento_geometrico(self):r=Logradouros(self.ruas(),'nome','EPSG:31983').consultar(QgsGeometry.fromWkt('POINT(50 2)'));self.assertEqual(len(r['cruzamentos']),1)
    def test_sem_rua_proxima(self):r=Logradouros(self.ruas(),'nome','EPSG:31983').consultar(QgsGeometry.fromWkt('POINT(5000 5000)'),10);self.assertFalse(r['candidatos'])
    def test_campo_ausente(self):
        with self.assertRaises(ValueError):Logradouros(self.ruas(),'inexistente','EPSG:31983')
    def test_crs_em_graus_bloqueado(self):
        with self.assertRaises(ValueError):Logradouros(self.ruas(),'nome','EPSG:4326')
    def test_externo_desabilitado(self):
        with self.assertRaises(PermissionError):Geocodificador().consultar(0,0,'EPSG:31983',{})
    def test_publico_nao_integrado(self):
        with self.assertRaisesRegex(ValueError,'pública'):Geocodificador().consultar(0,0,'EPSG:31983',dict(endpoint='https://nominatim.openstreetmap.org/reverse'),True)
    def test_credenciais_url_recusadas(self):
        with self.assertRaises(ValueError):Geocodificador().consultar(0,0,'EPSG:31983',dict(endpoint='https://identificador@example.invalid/reverse'),True)
    def test_atribuicao_sem_uri(self):r=sanear_fonte(self.ruas());self.assertNotIn('source',r);self.assertNotIn('uri',r)
    def test_enquadramento_minimo(self):r,a=rel.enquadrar([(QgsGeometry.fromWkt('POINT(0 0)'),'LOCAL')]);self.assertGreaterEqual(r.width(),100);self.assertFalse(a)
    def test_distantes_exigem_contexto(self):gs=[(QgsGeometry.fromWkt('POINT(0 0)'),'LOCAL'),(QgsGeometry.fromWkt('POINT(5000 0)'),'NO')];r,a=rel.enquadrar(gs,centro=gs[0][0]);self.assertTrue(a);self.assertLess(r.width(),1000)
    def test_enquadramento_invalido(self):
        with self.assertRaises(ValueError):rel.enquadrar([],margem=-1)
    def test_render_offline_sem_mudar_camada(self):
        with tempfile.TemporaryDirectory() as t:
            g=QgsGeometry.fromWkt('LINESTRING(0 0,100 0)');r=rel.renderizar([(g,'LINK')],'EPSG:31983',Path(t)/'map.png',{});self.assertTrue((Path(t)/'map.png').exists());self.assertEqual(g.asWkt(),'LineString (0 0, 100 0)');self.assertFalse(r['fontes'])
    def test_extensao_manual_escala_visivel(self):
        with tempfile.TemporaryDirectory() as t:
            r=rel.renderizar([(QgsGeometry.fromWkt('POINT(50 50)'),'LOCAL')],'EPSG:31983',Path(t)/'map.png',dict(extensao=[0,0,100,100]));a,b,c,d=r['extensao'];self.assertAlmostEqual((c-a)/(d-b),1.6,places=4)
    def test_base_sem_licenca_bloqueia(self):
        with tempfile.TemporaryDirectory() as t:
            with self.assertRaises(ValueError):rel.renderizar([(QgsGeometry.fromWkt('POINT(0 0)'),'LOCAL')],'EPSG:31983',Path(t)/'map.png',{},[self.ruas()])


class SelecaoRelatorioTest(unittest.TestCase):
    def setUp(self):self.cs=[dict(uid='a',origem='AUTOMATICA'),dict(uid='b',origem='MANUAL')];self.hs={('a','ESTRUTURAL'):dict(resultado='MANTER CONFIGURAÇÃO ATUAL')}
    def test_todos(self):self.assertEqual(len(rel.selecionar(self.cs,self.hs,'ESTRUTURAL')),2)
    def test_analisados(self):self.assertEqual(rel.selecionar(self.cs,self.hs,'ESTRUTURAL','ANALISADOS')[0][0]['uid'],'a')
    def test_pendentes(self):self.assertEqual(rel.selecionar(self.cs,self.hs,'ESTRUTURAL','PENDENTES')[0][0]['uid'],'b')
    def test_manual(self):self.assertEqual(rel.selecionar(self.cs,self.hs,'ESTRUTURAL','MANUAIS')[0][0]['uid'],'b')
    def test_filtros_ids_estaveis(self):self.assertEqual(len(rel.selecionar(self.cs,self.hs,'ESTRUTURAL','FILTROS',['a'])),1)
    def test_perspectiva_diferente(self):self.assertEqual(len(rel.selecionar(self.cs,self.hs,'TODAS_VALVULAS_ABERTAS','ANALISADOS')),0)


import test_decisoes_topologicas as fixture
from auditoria_topologica_watergems.diagnosticos import catalogar
from auditoria_topologica_watergems import operacional as op

class RelatorioPersistidoTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):fixture.DecisoesTest.setUpClass.__func__(cls)
    @classmethod
    def tearDownClass(cls):fixture.DecisoesTest.tearDownClass.__func__(cls)
    def setUp(self):fixture.DecisoesTest.setUp(self)
    def tearDown(self):fixture.DecisoesTest.tearDown(self)
    def test_snapshot_sem_gravar_na_fonte(self):
        antes=self.base.read_bytes();d=rel.snapshot(self.base,self.p,catalogar(self.p),'ESTRUTURAL',{});self.assertTrue(d['ocorrencias']);self.assertEqual(self.base.read_bytes(),antes)
    def test_relatorio_pdf_xlsx_mesma_selecao(self):
        from pypdf import PdfReader
        from openpyxl import load_workbook
        op.salvar(self.base,self.p,self.c,'ESTRUTURAL',resultado='MANTER CONFIGURAÇÃO ATUAL',observacao='Observação sintética')
        d=rel.gerar(self.base,self.p,catalogar(self.p),'ESTRUTURAL',self.dir/'relatorio',dict(escopo='SELECAO',ids=[self.c['uid']],imagens=False,xlsx=True))
        self.assertEqual(len(d['ocorrencias']),1);self.assertEqual(d['contagens']['MANTER CONFIGURAÇÃO ATUAL'],1);self.assertGreater(len(PdfReader(self.dir/'relatorio/INVESTIGACAO.pdf').pages),2)
        wb=load_workbook(self.dir/'relatorio/ACOMPANHAMENTO.xlsx');self.assertEqual(wb['OCORRENCIAS'].max_row,2);wb.close()
    def test_cancelamento_sem_publicacao(self):
        with self.assertRaises(InterruptedError):rel.gerar(self.base,self.p,catalogar(self.p),'ESTRUTURAL',self.dir/'cancelado',dict(imagens=False),cancelado=lambda:True)
        self.assertFalse((self.dir/'cancelado').exists())
    def test_destino_anterior_preservado(self):
        dest=self.dir/'antigo';dest.mkdir();(dest/'original').write_text('preservado')
        with self.assertRaises(FileExistsError):rel.gerar(self.base,self.p,catalogar(self.p),'ESTRUTURAL',dest,{})
        self.assertEqual((dest/'original').read_text(),'preservado')
    def test_falha_pdf_nao_publica(self):
        with patch.object(rel,'gerar_pdf',side_effect=OSError('Permissão recusada')):
            with self.assertRaises(OSError):rel.gerar(self.base,self.p,catalogar(self.p),'ESTRUTURAL',self.dir/'falha',dict(imagens=False))
        self.assertFalse((self.dir/'falha').exists());self.assertFalse(list(self.dir.glob('.relatorio_*')))
    def test_historico_preserva_revisoes(self):
        a=op.salvar(self.base,self.p,self.c,'ESTRUTURAL',resultado='MANTER CONFIGURAÇÃO ATUAL');op.salvar(self.base,self.p,self.c,'ESTRUTURAL',resultado='CORREÇÃO NECESSÁRIA',esperada=a)
        d=rel.snapshot(self.base,self.p,[self.c],'ESTRUTURAL',dict(historico=True));self.assertEqual(len(d['ocorrencias'][0]['historico']),2)
    def test_renomear_sem_regravar_gpkg(self):
        cat=Catalogo(self.dir);cat.registrar(self.base);r=next(r for r in cat.listar() if r['tipo']=='INTEGRADA');antes=self.base.read_bytes();gc.alterar(cat,r,'RENOMEAR','Investigação de teste');self.assertEqual(self.base.read_bytes(),antes)
    def test_mover_e_recuperar_reversivel(self):
        cat=Catalogo(self.dir);cat.registrar(self.base);r=next(r for r in cat.listar() if r['tipo']=='INTEGRADA');antes=self.base.read_bytes();d=gc.recuperar_arquivo(cat,r,confirmado=True);self.assertFalse(self.base.exists());m=Path(d['recuperacao']).with_suffix('.recuperacao.json');gc.restaurar_arquivo(cat,m);self.assertEqual(self.base.read_bytes(),antes)
    def test_dependencia_outra_investigacao(self):
        cat=Catalogo(self.dir);cat.registrar(self.base);r=next(r for r in cat.listar() if r['tipo']=='INTEGRADA')
        with sqlite3.connect(self.dir/'referencia.gpkg') as c:
            c.execute('CREATE TABLE IG_EXECUCOES(ig_uid TEXT,data_utc TEXT,meta TEXT)');m=copy.deepcopy(self.p['resultado']['meta']);m['geo']['base']=str(self.base);m['uid']='outra';c.execute('INSERT INTO IG_EXECUCOES VALUES(?,?,?)',('outra','hoje',json.dumps(m)))
        cat.registrar(self.dir/'referencia.gpkg');self.assertTrue(any('depende' in b for b in gc.dependencias(cat,r)['bloqueios']))
    def test_dialogos_acessiveis(self):
        j=fixture.DecisoesTest.painel(self)
        from auditoria_topologica_watergems.relatorio_ui import RelatorioDialog
        d=RelatorioDialog(j);self.assertFalse(d.externo.isChecked());self.assertEqual(d.config()['escopo'],'TODOS');d.close();d.deleteLater()


def dispositivo_tipo(tipo):
    def testar(self):
        p,h=exemplo();g=p['topologia']['estrutural']['grafo'];el=p['topologia']['modelo']['elementos'];el[2].update(familia='DISPOSITIVO',tipo=tipo)
        for a,b,k,d in list(g.edges(keys=True,data=True)):
            if d['elemento']==102:g.remove_edge(a,b,k)
        g.add_edge((2,'S'),(3,'E'),elemento=102,natureza='TUBULACAO');g.add_edge((3,'E'),(2,'S'),elemento=102,natureza='TUBULACAO');g.add_edge((2,'E'),(2,'S'),elemento=2,natureza='DISPOSITIVO')
        q=calcular(p,'ESTRUTURAL',[1],[3],comprimentos={101:10,102:20});h.update(vinculos=q['elementos'],dados=dict(criterios=q));r=ex.preparar(p,h);self.assertIn(tipo,r['contagens']);self.assertEqual([e['Elemento'] for e in r['percursos']],[101,2,102])
    return testar
for tipo in ('PRV','TCV','FCV','PBV','CheckValve','StandardPump'):
    setattr(ExportacaoCompletaTest,'test_percurso_com_'+tipo,dispositivo_tipo(tipo))

from qgis.PyQt.QtCore import QObject,pyqtSignal,QByteArray,QTimer
from qgis.PyQt.QtNetwork import QNetworkReply
from qgis.core import QgsNetworkAccessManager
class RespostaTeste(QObject):
    finished=pyqtSignal()
    def __init__(self,body=b'{"display_name":"Rua sintetica"}',status=200,erro=QNetworkReply.NoError):super().__init__();self.body=body;self.status=status;self.err=erro
    def isFinished(self):return False
    def error(self):return self.err
    def attribute(self,k):return self.status
    def bytesAvailable(self):return len(self.body)
    def readAll(self):return QByteArray(self.body)
    def abort(self):self.err=QNetworkReply.OperationCanceledError;self.finished.emit()
class RedeTeste:
    def __init__(self,r):self.r=r;self.reqs=[]
    def get(self,req):self.reqs.append(req);QTimer.singleShot(0,self.r.finished.emit);return self.r
class ServicoLocalizacaoTest(unittest.TestCase):
    def conf(self):return dict(endpoint='https://cartografia.example.invalid/reverse',termos='Contrato sintético permite consulta unitária',atribuicao='Fonte de teste',custos='Sem custo no mock',armazenamento='Permite cache local',cache_permitido=True)
    def test_envia_somente_coordenadas(self):
        from urllib.parse import parse_qs,urlsplit
        rede=RedeTeste(RespostaTeste())
        with patch.object(QgsNetworkAccessManager,'instance',return_value=rede):r=Geocodificador().consultar(500000,7400000,'EPSG:31983',self.conf(),True)
        self.assertEqual(set(parse_qs(urlsplit(rede.reqs[0].url().toString()).query)),{'lat','lon','format'});self.assertEqual(r['fonte'],'cartografia.example.invalid')
    def test_cache_autorizado_nao_repete(self):
        rede=RedeTeste(RespostaTeste());g=Geocodificador()
        with patch.object(QgsNetworkAccessManager,'instance',return_value=rede):
            a=g.consultar(500000,7400000,'EPSG:31983',self.conf(),True);b=g.consultar(500000,7400000,'EPSG:31983',self.conf(),True)
        self.assertEqual(a,b);self.assertEqual(len(rede.reqs),1)
    def test_intervalo_respeitado(self):
        rede=RedeTeste(RespostaTeste());g=Geocodificador()
        with patch.object(QgsNetworkAccessManager,'instance',return_value=rede):
            g.consultar(500000,7400000,'EPSG:31983',self.conf(),True)
            with self.assertRaises(ValueError):g.consultar(500001,7400000,'EPSG:31983',self.conf(),True)
    def test_licenca_ausente_sem_rede(self):
        c=self.conf();del c['termos']
        with self.assertRaises(ValueError):Geocodificador().consultar(500000,7400000,'EPSG:31983',c,True)

def falha_servico(body,status,erro):
    def testar(self):
        rede=RedeTeste(RespostaTeste(body,status,erro))
        with patch.object(QgsNetworkAccessManager,'instance',return_value=rede):
            with self.assertRaises((ValueError,RuntimeError)):Geocodificador().consultar(500000,7400000,'EPSG:31983',self.conf(),True)
    return testar
for nome,body,status,erro in [('http_invalido',b'{}',500,QNetworkReply.NoError),('json_invalido',b'not json',200,QNetworkReply.NoError),('autenticacao',b'{}',401,QNetworkReply.AuthenticationRequiredError),('timeout',b'{}',200,QNetworkReply.TimeoutError),('rate_limit',b'{}',429,QNetworkReply.NoError),('sem_endereco',b'{}',200,QNetworkReply.NoError)]:
    setattr(ServicoLocalizacaoTest,'test_'+nome,falha_servico(body,status,erro))

class ReferenciaCadastralTest(unittest.TestCase):
    def test_zeros_preservados(self):
        from auditoria_topologica_watergems.localizacao_relatorio import referencia_cadastral
        p=dict(dossie=dict(tubulacoes=[dict(uid='t',origem='origem',atributos={'rua':'001'})]));r=referencia_cadastral(p,dict(gis_uids=['t']),dict(origem='origem',campo='rua'));self.assertEqual(r['valores'],['001'])
    def test_nao_presume_nome(self):
        from auditoria_topologica_watergems.localizacao_relatorio import referencia_cadastral
        p=dict(dossie=dict(tubulacoes=[dict(uid='t',origem='origem',atributos={'NOME':'Qualquer'})]));self.assertIsNone(referencia_cadastral(p,dict(gis_uids=['t']),{}))

if __name__=='__main__':unittest.main()
