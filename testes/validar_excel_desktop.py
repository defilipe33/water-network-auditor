"""Teste opcional com uma instância própria e oculta do Excel desktop (COM).

Abre normalmente (CorruptLoad=0), verifica tabelas/filtros e salva uma cópia
nova para reabertura. Nunca salva sobre a entrada nem usa instâncias existentes.
"""
import argparse
import hashlib
import json
from pathlib import Path
import sys


def validar_excel(entrada, copia):
    import pythoncom
    from win32com.client import DispatchEx
    from openpyxl import load_workbook
    entrada=Path(entrada).resolve(strict=True)
    copia=Path(copia).resolve()
    if entrada==copia or copia.exists():
        raise ValueError('A cópia de reabertura precisa ter um nome novo')
    original=hashlib.sha256(entrada.read_bytes()).hexdigest()
    esperado={}
    wb=load_workbook(entrada)
    for ws in wb:
        esperado[ws.title]=dict(tabelas=[dict(nome=t.name,ref=t.ref,colunas=len(t.tableColumns))
                                        for t in ws.tables.values()],
                                valores=list(ws.values),
                                outline={n:dict(nivel=r.outlineLevel,oculta=bool(r.hidden),recolhida=bool(r.collapsed))
                                    for n,r in ws.row_dimensions.items() if r.outlineLevel or r.collapsed})
    wb.close()
    pythoncom.CoInitialize()
    app=None
    logs=[]
    try:
        app=DispatchEx('Excel.Application')
        app.Visible=False
        app.DisplayAlerts=False
        app.EnableEvents=False
        app.AskToUpdateLinks=False
        app.AutomationSecurity=3
        versao=str(app.Version)
        build=str(app.Build)
        def abrir(caminho):
            m=pythoncom.Missing
            return app.Workbooks.Open(str(caminho),0,True,m,m,m,True,m,m,False,False,m,False,True,0)
        for fase,caminho in [('ABERTURA_NORMAL',entrada),('REABERTURA_COPIA_EXCEL',copia)]:
            livro=abrir(caminho)
            try:
                abas=[]
                for nome,dados in esperado.items():
                    ws=livro.Worksheets.Item(nome)
                    if ws.ListObjects.Count!=len(dados['tabelas']):
                        raise AssertionError(f'{fase}/{nome}: Excel retirou ou alterou tabelas')
                    for t in dados['tabelas']:
                        tabela=ws.ListObjects.Item(t['nome'])
                        ref=str(tabela.Range.Address).replace('$','')
                        if ref!=t['ref'] or tabela.ListColumns.Count!=t['colunas']:
                            raise AssertionError(f'{fase}/{nome}: intervalo ou colunas alterados')
                        if not tabela.ShowAutoFilter:
                            raise AssertionError(f'{fase}/{nome}: filtro da tabela ausente')
                        primeiro=tabela.DataBodyRange.Cells.Item(1,1).Value2
                        tabela.Range.AutoFilter(1,primeiro)
                        if not tabela.AutoFilter.Filters.Item(1).On:
                            raise AssertionError(f'{fase}/{nome}: aplicação do filtro não funcionou')
                        tabela.AutoFilter.ShowAllData()
                    ws.Activate()
                    janela=livro.Windows.Item(1)
                    if not janela.FreezePanes or janela.SplitRow!=1:
                        raise AssertionError(f'{fase}/{nome}: cabeçalho congelado perdido')
                    if dados['outline']:
                        if ws.Outline.SummaryRow!=0:
                            raise AssertionError('Excel alterou a posição dos resumos do outline')
                        for numero,esperada in dados['outline'].items():
                            linha=ws.Rows.Item(numero)
                            if linha.OutlineLevel!=esperada['nivel']+1 or bool(linha.Hidden)!=esperada['oculta']:
                                raise AssertionError('Excel alterou níveis ou recolhimento das linhas')
                        principais=[n for n,r in dados['outline'].items() if r['recolhida'] and not r['nivel']]
                        primeiro=principais[0]
                        proximo=min((n for n in principais if n>primeiro),default=ws.UsedRange.Rows.Count+1)
                        ws.Rows.Item(primeiro).ShowDetail=True
                        if any(ws.Rows.Item(n).Hidden for n in range(primeiro+1,proximo)):
                            raise AssertionError('Expansão individual não mostrou os detalhes')
                        if len(principais)>1 and not ws.Rows.Item(principais[1]+1).Hidden:
                            raise AssertionError('Expandir um local alterou outro local')
                        ws.Rows.Item(primeiro).ShowDetail=False
                        if any(not ws.Rows.Item(n).Hidden for n in range(primeiro+1,proximo)):
                            raise AssertionError('Recolhimento individual não ocultou os detalhes')
                    abas.append(dict(nome=nome,tabelas=ws.ListObjects.Count,
                                     intervalos=[t['ref'] for t in dados['tabelas']],
                                     outline_validado=bool(dados['outline'])))
                logs.append(dict(fase=fase,abas=abas))
                if fase=='ABERTURA_NORMAL':
                    livro.SaveCopyAs(str(copia))
            finally:
                livro.Close(False)
        if hashlib.sha256(entrada.read_bytes()).hexdigest()!=original:
            raise AssertionError('Arquivo de entrada foi alterado pelo teste')
    finally:
        if app is not None:
            app.Quit()
        app=None
        pythoncom.CoUninitialize()
    # Comparação integral das células após a gravação pelo Excel real.
    gravado=load_workbook(copia)
    for ws in gravado:
        valores=list(ws.values)
        inicial=esperado[ws.title]['valores']
        def normalizar(linhas):
            return [tuple(None if v=='' else v for v in linha) for linha in linhas]
        if normalizar(valores)!=normalizar(inicial):
            raise AssertionError(f'Dados alterados após salvar/reabrir: {ws.title}')
    gravado.close()
    return dict(excel=versao,build=build,modo_abertura='Normal (0), sem recuperação solicitada',
                validacoes=logs,celulas_preservadas=True,entrada_intacta=True,copia=str(copia))


if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('arquivo')
    p.add_argument('copia')
    p.add_argument('--controle-anterior',help='Arquivo anterior, somente leitura, para reproduzir o defeito')
    p.add_argument('--log',help='Caminho novo para registrar a validação em JSON')
    args=p.parse_args()
    try:
        controle=None
        if args.controle_anterior:
            try:
                controle=validar_excel(args.controle_anterior,
                    Path(args.copia).with_name('CONTROLE_ANTERIOR_excel_roundtrip.xlsx'))
            except Exception as erro:
                controle=dict(sucesso=False,erro=str(erro))
        resultado=dict(controle_anterior=controle,corrigido=validar_excel(args.arquivo,args.copia))
        if args.log:
            with Path(args.log).open('x',encoding='utf-8') as arquivo:
                json.dump(resultado,arquivo,ensure_ascii=False,indent=2)
        print(json.dumps(resultado,ensure_ascii=False,indent=2))
    except Exception as erro:
        import traceback
        traceback.print_exc()
        print(json.dumps(dict(sucesso=False,erro=str(erro)),ensure_ascii=False))
        sys.exit(1)
