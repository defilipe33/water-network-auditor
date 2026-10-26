"""Confere XLSX/JSON do ensaio e, opcionalmente, preservação de um ensaio anterior."""
import argparse
from collections import Counter
import json
import math
from pathlib import Path

from openpyxl import load_workbook
from validar_estrutura_xlsx import validar_xlsx


def validar(caminho, anterior=None):
    caminho=Path(caminho)
    dados=json.loads(caminho.with_name(caminho.stem+'_tecnico.json').read_text(encoding='utf-8'))
    estrutura=validar_xlsx(caminho)
    assert estrutura['valido'],estrutura['erros']
    eventos=dados['ocorrencias']
    analise=dados['investigacao']
    locais=analise['locais']
    locais_exibidos=dados.get('dossie',{}).get('locais',locais)
    vinculos=analise['vinculos']
    uids={e['uid'] for e in eventos}
    assert len(uids)==len(eventos)
    if analise['opcoes']['habilitado']:
        assert Counter(v['evento_uid'] for v in vinculos)==Counter(uids)
        assert len({l['uid'] for l in locais})==len(locais)
        assert {v['local_uid'] for v in vinculos}=={l['uid'] for l in locais}
        for local in locais:
            membros={v['evento_uid'] for v in vinculos if v['local_uid']==local['uid']}
            assert membros==set(local['eventos'])
            assert len(membros)==local['registro']['N_EVENTOS']
    else:
        assert not locais and not vinculos
    wb=load_workbook(caminho)
    try:
        def conferir(nome, registros):
            reais=list(wb[nome].iter_rows(min_row=2,values_only=True))
            campos=[c.value for c in wb[nome][1]]
            esperados=[tuple(r[c] for c in campos) for r in registros]
            assert len(reais)==len(esperados),nome
            for real,esperado in zip(reais,esperados):
                for encontrado,valor in zip(real,esperado):
                    if isinstance(valor,float):
                        # XLSX limita a precisão decimal; 10 nanômetros nas coordenadas.
                        assert isinstance(encontrado,(int,float)) and math.isclose(
                            encontrado,valor,rel_tol=0,abs_tol=1e-8),(nome,encontrado,valor)
                    else:
                        assert encontrado==(None if valor=='' else valor),(nome,encontrado,valor)
        for modo,nome in [('TRANSICAO','TRANSICOES_SETORES'),('PERIMETRO','TRAVESSIAS_PERIMETROS')]:
            conferir(nome,[e['registro'] for e in eventos if e['tecnico']['modo']==modo])
        conferir('PONTOS_INVESTIGACAO',[l['registro'] for l in locais_exibidos])
        conferir('PARES_SETORES',analise['pares'])
        conferir('SAIDAS_EXTERNAS',[s['registro'] for s in analise['saidas']])
        resumo={r[1]:r[2] for r in wb['RESUMO'].iter_rows(min_row=2,values_only=True)
                if r[0] in ('TOTAL','INVESTIGACAO')}
        assert resumo['Ocorrências únicas']==len(eventos)
        assert resumo['Locais de investigação']==len(locais)
        assert resumo['Travessias região/exterior inequívocas']==len(analise['saidas'])
        assert resumo['Pares inequívocos distintos']==sum(p['LEITURA']=='INEQUIVOCA' for p in analise['pares'])
        assert sum(l['registro']['N_EVENTOS'] for l in locais)==(len(eventos) if locais else 0)
    finally:
        wb.close()
    preservados=None
    tecnicos_preservados=None
    if anterior:
        antes=json.loads(Path(anterior).read_text(encoding='utf-8'))
        assert antes['parametros']==dados['parametros'],'Parâmetros diferentes: comparação de regressão inadequada'
        assert antes['fontes']==dados['fontes'],'Entradas/configurações diferentes'
        preservados=antes['ocorrencias']==eventos
        def bruto(ocorrencias):
            saida={}
            for e in ocorrencias:
                t=e['tecnico'].copy()
                t.pop('apresentacao',None)
                saida[e['uid']]=t
            return saida
        tecnicos_preservados=bruto(antes['ocorrencias'])==bruto(eventos)
        assert tecnicos_preservados,'Conteúdo geométrico/técnico original divergente'
    return dict(valido=True,estrutura=estrutura,contadores=analise['contadores'],
        eventos_anteriores_integralmente_preservados=preservados,
        dados_geometricos_e_tecnicos_anteriores_preservados=tecnicos_preservados,
        dados_abas_iguais_ao_json=True,vinculos_integros=True,
        distribuicao_locais_por_quantidade=dict(Counter(l['registro']['N_EVENTOS'] for l in locais)),
        exemplos=[l['registro'] for l in locais if l['registro']['N_EVENTOS']>1][:5])


if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('arquivo')
    p.add_argument('--anterior',help='JSON técnico do mesmo conjunto de entradas e parâmetros')
    p.add_argument('--log',help='Novo arquivo JSON para evidências')
    args=p.parse_args()
    evidencia=validar(args.arquivo,args.anterior)
    texto=json.dumps(evidencia,ensure_ascii=False,indent=2)
    if args.log:
        with Path(args.log).open('x',encoding='utf-8') as arq:
            arq.write(texto)
    print(texto)
