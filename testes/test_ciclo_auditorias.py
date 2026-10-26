"""Ciclo de vida real PyQGIS: versões, regeneração, reabertura e falhas."""
from copy import deepcopy
from pathlib import Path
import json
import tempfile
import unittest
from unittest.mock import patch
from test_auditoria_fronteiras import APP,camada,quadrado
from qgis.core import QgsProject,QgsVectorLayer
from auditoria_fronteiras import auditar,CamadaEntrada,Parametros
from auditoria_fronteiras.investigacao import analisar_investigacao,OpcoesLocais
from auditoria_fronteiras.inspecao import construir_dossie,restaurar,texto
from auditoria_fronteiras.ciclo import (inicializar,regerar,perfil,comparar_locais,
    comparar_configuracoes,configuracao_proposta,assinatura_camada)
from auditoria_fronteiras.versoes import (preparar_regeneracao,gravar_versao,listar,
    validar_base,backup_verificado,abrir_ro,salvar_preferencias)
from auditoria_fronteiras.persistencia import salvar_execucao,carregar_execucao,salvar_revisao,historico,conectar
from auditoria_fronteiras.catalogo import publicar,carregadas,lembrar,recentes,resolver_fontes
from auditoria_fronteiras.interface import AuditoriaFronteirasDialog


def preparar():
    rede=camada('LineString',['LINESTRING(-2 5,12 5)'],nome='rede')
    p=camada('Polygon',[quadrado()],nome='externo')
    q=camada('Polygon',[quadrado(2,2,6)],nome='interno')
    redes=[CamadaEntrada(rede,'id')];limites=[CamadaEntrada(c,'id','id') for c in (p,q)]
    r=auditar(redes,limites,Parametros())
    a=analisar_investigacao(r,OpcoesLocais(True,5,5))
    d=inicializar(construir_dossie(r,a),'Teste','Município livre')
    d['configuracao']=configuracao_proposta(redes,limites,r.parametros,a.opcoes)
    return d,redes,limites


class CicloTest(unittest.TestCase):
    def setUp(self):
        for metodo in ('confirmar_processamento','resolver_pendencias'):
            p=patch.object(AuditoriaFronteirasDialog,metodo,return_value=True)
            p.start();self.addCleanup(p.stop)
        self.tmp=tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        self.pasta=Path(self.tmp.name);self.base=self.pasta/'auditoria.gpkg'
        self.d,self.redes,self.limites=preparar()
        self.uid=self.d['execucao_uid']
        salvar_execucao(self.base,self.d)

    def tearDown(self):
        self.tmp.cleanup()

    def preparar(self,distancia=0,habilitado=True):
        return preparar_regeneracao(self.base,self.uid,OpcoesLocais(habilitado,distancia,distancia))

    def gravar(self,p,modo='NOVA',**kw):
        return gravar_versao(self.base,self.uid,p['novo'],modo,token_origem=p['token_origem'],**kw)

    def revisar(self):
        loc=self.d['locais'][0]['uid']
        salvar_revisao(self.base,self.uid,loc,'VERIFICADO','observação preservada','Engenheiro')
        return loc

    def test_regerar_distancias_sem_nucleo_preserva_todos_eventos(self):
        with patch('auditoria_fronteiras.nucleo.auditar',side_effect=AssertionError('núcleo proibido')):
            for distancia,quantidade in [(5,2),(0,4),(1,4),(2,2),(10,1)]:
                with self.subTest(distancia=distancia):
                    p=self.preparar(distancia)
                    self.assertEqual(len(p['novo']['locais']),quantidade)
                    for c in ('eventos','contexto','tubulacoes','poligonos','evento_rede','evento_poligono'):
                        self.assertEqual(texto(p['novo'][c]),texto(self.d[c]))

    def test_desativar_e_reativar_sem_apagar_eventos(self):
        p=self.preparar(5,False)
        self.assertEqual(p['novo']['locais'],[]);self.assertEqual(p['novo']['local_evento'],[])
        novo,_=regerar(p['novo'],OpcoesLocais(True,5,5))
        self.assertEqual(len(novo['locais']),2)
        self.assertEqual(novo['eventos'],self.d['eventos'])

    def test_divisao_e_fusao_efetivas(self):
        p=self.preparar(0)
        self.assertEqual((p['comparacao']['divididos'],p['comparacao']['fundidos']),(2,0))
        p=self.preparar(10)
        self.assertEqual((p['comparacao']['divididos'],p['comparacao']['fundidos']),(0,1))

    def test_exatas_e_parciais_por_composicao(self):
        self.assertEqual(self.preparar(5)['comparacao']['mantidos'],2)
        self.assertEqual(self.preparar(0)['comparacao']['mantidos'],0)
        self.assertTrue(all(r['tipo']=='PARCIAL' for r in self.preparar(0)['comparacao']['correspondencias']))

    def test_nova_versao_padrao_nao_herda_revisao(self):
        loc=self.revisar();novo=self.gravar(self.preparar(5))['dossie']
        self.assertTrue(all(l['revisao']['situacao']=='PENDENTE' for l in novo['locais']))
        self.assertEqual(historico(self.base,self.uid,loc)[0]['observacoes'],'observação preservada')
        self.assertEqual(len(listar(self.base)),2)

    def test_reconciliar_equivalentes_preserva_revisao(self):
        loc=self.revisar();novo=self.gravar(self.preparar(5),'RECONCILIAR')['dossie']
        l=next(l for l in novo['locais'] if l['uid']==loc)
        self.assertEqual(l['revisao']['situacao'],'VERIFICADO')
        self.assertEqual(len(historico(self.base,self.uid,loc)),1)
        self.assertTrue(novo['linhagem_locais'])

    def test_divisao_bloqueia_heranca_de_verificado(self):
        self.revisar();p=self.preparar(0)
        self.assertEqual(p['comparacao']['revisoes_afetadas'],1)
        novo=self.gravar(p,'RECONCILIAR')['dossie']
        self.assertTrue(all(l['revisao']['situacao']=='PENDENTE' for l in novo['locais']))
        self.assertTrue(any(r['tipo']=='PARCIAL' and r['decisao']=='PENDENTE' for r in novo['reconciliacoes']))

    def test_fusao_revisoes_conflitantes_nao_herdadas(self):
        self.revisar();loc=self.d['locais'][1]['uid']
        salvar_revisao(self.base,self.uid,loc,'NECESSITA_CONFIRMACAO','outra decisão')
        novo=self.gravar(self.preparar(10),'RECONCILIAR')['dossie']
        self.assertEqual(novo['locais'][0]['revisao']['situacao'],'PENDENTE')
        self.assertEqual(len(novo['linhagem_locais']),2)

    def test_reabrir_versoes_antigas_e_novas(self):
        loc=self.revisar();novo=self.gravar(self.preparar(0))['dossie']
        velho=carregar_execucao(self.base,self.uid)
        self.assertEqual(len(velho['locais']),2);self.assertEqual(len(novo['locais']),4)
        self.assertEqual(next(l for l in velho['locais'] if l['uid']==loc)['revisao']['situacao'],'VERIFICADO')
        for d in (velho,carregar_execucao(self.base,novo['execucao_uid'])):
            r,a=restaurar(d);self.assertEqual(len(r.ocorrencias),4)

    def test_substituir_exige_confirmacao(self):
        with self.assertRaisesRegex(ValueError,'confirmação'):
            self.gravar(self.preparar(),'SUBSTITUIR')

    def test_substituir_backup_recuperavel_e_anterior_preservada(self):
        self.revisar();s=self.gravar(self.preparar(),'SUBSTITUIR',confirmado=True)
        self.assertTrue(validar_base(s['backup']['backup']))
        self.assertEqual(len(listar(s['backup']['backup'])),1)
        self.assertFalse(carregar_execucao(self.base,self.uid)['controle_versao']['ativa'])
        self.assertTrue(carregar_execucao(self.base,s['dossie']['execucao_uid'])['controle_versao']['ativa'])
        self.assertEqual(len(carregar_execucao(s['backup']['backup'],self.uid)['locais']),2)

    def test_falha_antes_commit_reverte_toda_versao(self):
        p=self.preparar()
        def falhar(etapa):
            if etapa=='ANTES_COMMIT':
                raise RuntimeError('falha simulada')
        with self.assertRaisesRegex(RuntimeError,'simulada'):
            self.gravar(p,interromper=falhar)
        self.assertEqual(len(listar(self.base)),1)
        self.assertTrue(validar_base(self.base))

    def test_falha_substituicao_preserva_backup_e_versao_ativa(self):
        p=self.preparar()
        with self.assertRaises(RuntimeError):
            self.gravar(p,'SUBSTITUIR',confirmado=True,interromper=lambda etapa: (_ for _ in ()).throw(RuntimeError('falha')))
        self.assertTrue(carregar_execucao(self.base,self.uid)['controle_versao']['ativa'])
        self.assertEqual(len(list((self.pasta/'backups').glob('*.gpkg'))),1)

    def test_revisao_concorrente_bloqueia_comparacao_obsoleta(self):
        p=self.preparar();self.revisar()
        with self.assertRaisesRegex(ValueError,'mudaram'):
            self.gravar(p)

    def test_cancelamento_regeneracao_nao_grava(self):
        with self.assertRaises(Exception):
            preparar_regeneracao(self.base,self.uid,OpcoesLocais(True),lambda:True)
        self.assertEqual(len(listar(self.base)),1)

    def test_evento_modificado_impede_regeneracao(self):
        p=self.preparar();p['novo']['eventos'][0]['registro']['AREA']='alterada'
        with self.assertRaises(ValueError):
            self.gravar(p)

    def test_backup_completo_tabelas_historico(self):
        self.revisar();b=self.pasta/'backup.gpkg'
        ev=backup_verificado(self.base,b,self.uid)
        self.assertTrue(ev['estrutura_dados_iguais'])
        self.assertEqual(len(historico(b,self.uid,self.d['locais'][0]['uid'])),1)
        with self.assertRaises(ValueError):
            backup_verificado(self.base,b)

    def test_relacionamentos_de_cada_versao(self):
        novo=self.gravar(self.preparar())['dossie']
        self.assertEqual(len(novo['local_evento']),4)
        self.assertEqual(len({r['evento_uid'] for r in novo['local_evento']}),4)
        with abrir_ro(self.base) as con:
            self.assertFalse(con.execute('PRAGMA foreign_key_check').fetchall())
            self.assertEqual(con.execute('SELECT count(*) FROM AUDITORIA_LINHAGEM_EVENTOS').fetchone()[0],4)

    def test_perfil_e_preferencias_recuperados(self):
        salvar_preferencias(self.base,self.uid,{'formatos':['SHP'],'destino':'pasta anterior'})
        d=carregar_execucao(self.base,self.uid)
        self.assertEqual(perfil(d)['exportacao']['formatos'],['SHP'])
        self.assertEqual(perfil(d)['fontes'],perfil(self.d)['fontes'])

    def test_configuracoes_distinguem_geometria_e_agrupamento(self):
        p=perfil(self.d);q=deepcopy(p);q['agrupamento']['distancia']=0
        self.assertTrue(comparar_configuracoes(p,q)['recomendar_regerar'])
        q['geometria']['tolerancia_geometrica']=.1
        self.assertTrue(comparar_configuracoes(p,q)['requer_nucleo'])

    def test_alterar_conteudo_mesma_uri_exige_nucleo(self):
        p=perfil(self.d);q=deepcopy(p)
        origem=next(iter(q['assinaturas_fontes']));q['assinaturas_fontes'][origem]='modificado'
        self.assertTrue(comparar_configuracoes(p,q)['requer_nucleo'])

    def test_reprocessar_entradas_e_tolerancias(self):
        r=auditar(self.redes,self.limites[:1],Parametros(tolerancia_geometrica=.05))
        a=analisar_investigacao(r,OpcoesLocais(True));d=inicializar(construir_dossie(r,a),origem=self.d)
        d['configuracao']=configuracao_proposta(self.redes,self.limites[:1],r.parametros,a.opcoes)
        p=self.preparar();p['novo']=d
        novo=self.gravar(p)['dossie']
        self.assertEqual(len(novo['eventos']),2)
        self.assertEqual(novo['contexto']['parametros']['tolerancia_geometrica'],.05)
        self.assertEqual(len(carregar_execucao(self.base,self.uid)['eventos']),4)

    def test_nao_herdar_quando_contexto_muda(self):
        self.revisar();p=self.preparar(5)
        p['novo']['configuracao']['geometria']['tolerancia_geometrica']=.1
        self.assertEqual(comparar_locais(self.d,p['novo'])['mantidos'],0)
        novo=self.gravar(p,'RECONCILIAR')['dossie']
        self.assertTrue(all(l['revisao']['situacao']=='PENDENTE' for l in novo['locais']))

    def test_reabrir_sem_duplicar_camadas_e_reutilizar_eventos(self):
        projeto=QgsProject()
        _,_,cs=publicar(self.d,self.base,projeto)
        ids=[c.id() for c in cs];quantidade=len(projeto.mapLayers())
        _,_,cs2=publicar(carregar_execucao(self.base,self.uid),self.base,projeto)
        self.assertEqual([c.id() for c in cs2],ids)
        self.assertEqual(len(projeto.mapLayers()),quantidade)
        novo=self.gravar(self.preparar())['dossie']
        _,_,cs3=publicar(novo,self.base,projeto)
        self.assertEqual(cs3[0].id(),ids[0])
        for c in projeto.mapLayers().values():
            c.setName('resultado renomeado')
        itens,avisos=carregadas(projeto)
        self.assertEqual({x['execucao_uid'] for x in itens},{self.uid,novo['execucao_uid']})
        projeto.clear()

    def test_camadas_gpkg_carregadas_sem_nome_fixo(self):
        projeto=QgsProject();c=QgsVectorLayer(str(self.base)+'|layername=EVENTOS_FRONTEIRAS','nome livre','ogr')
        projeto.addMapLayer(c)
        itens,_=carregadas(projeto)
        self.assertEqual(itens[0]['execucao_uid'],self.uid)
        projeto.clear()

    def test_catalogo_recent_e_realocacao_por_identidade(self):
        arquivo=self.pasta/'recentes.json';lembrar(self.base,self.uid,arquivo)
        lembrar(self.base,self.uid,arquivo)
        self.assertEqual(len(recentes(arquivo)),1)
        b=self.pasta/'movida.gpkg';backup_verificado(self.base,b)
        lembrar(b,self.uid,arquivo)
        self.assertEqual(recentes(arquivo)[0]['base'],str(b.resolve()))

    def test_resolver_fontes_e_recuperar_formulario_sem_nomes(self):
        projeto=QgsProject()
        for c in self.redes+self.limites:
            c.camada.setName('nome alterado');projeto.addMapLayer(c.camada)
        fontes,_=resolver_fontes(perfil(self.d),projeto)
        self.assertEqual(len(fontes),3)
        j=AuditoriaFronteirasDialog(projeto=projeto)
        j.modalidade.setCurrentIndex(3);j.restaurar_configuracao(self.d,True)
        redes,limites,p=j.configuracoes()
        self.assertEqual(len(redes),1);self.assertEqual(len(limites),2)
        self.assertTrue(all(c.chave_origem for c in redes+limites))
        self.assertTrue(all(c.campo_nome=='id' for c in limites))
        self.assertEqual(j.distancia_locais.value(),5)
        j.close();projeto.clear()

    def test_abrir_formulario_nao_chama_nucleo_nem_agrupador(self):
        projeto=QgsProject();j=AuditoriaFronteirasDialog(projeto=projeto)
        with patch('auditoria_fronteiras.interface.auditar',side_effect=AssertionError()),\
             patch('auditoria_fronteiras.investigacao.analisar_investigacao',side_effect=AssertionError()),\
             patch('auditoria_fronteiras.fluxos.lembrar'):
            j.abrir_auditoria(self.base,self.uid)
            j.abrir_auditoria(self.base,self.uid)
        self.assertEqual(len(projeto.mapLayers()),4)
        j.close();projeto.clear()

    def test_integridade_rejeita_gpkg_incompleto(self):
        p=self.pasta/'incompleto.gpkg';p.write_bytes(b'')
        with self.assertRaises(ValueError):
            validar_base(p)

    def test_exportacoes_apos_regenerar(self):
        from auditoria_fronteiras.exportacao import exportar_shp,validar_gpkg
        from auditoria_fronteiras.relatorio import exportar_xlsx
        from validar_estrutura_xlsx import validar_xlsx
        from openpyxl import load_workbook
        novo=self.gravar(self.preparar())['dossie'];r,a=restaurar(novo)
        self.assertTrue(validar_gpkg(self.base,novo))
        exportar_shp(novo,self.pasta/'SHP')
        x=self.pasta/'relatorio.xlsx';exportar_xlsx(r,x,analise=a,dossie=novo)
        validar_xlsx(x)
        wb=load_workbook(x)
        self.assertEqual(wb['PONTOS_INVESTIGACAO'].max_row,5)
        ws=wb['PARAMETROS'];linha=next(c.row for c in ws['C'] if c.value=='tolerancia_geometrica')
        self.assertEqual(ws.cell(linha,4).value,.01)
        self.assertEqual(ws.cell(linha,4).number_format,'0.############')
        self.assertIn('DETALHAMENTO_LOCAIS',wb.sheetnames);wb.close()

    def test_regenerar_interface_nunca_chama_nucleo(self):
        projeto=QgsProject();j=AuditoriaFronteirasDialog(projeto=projeto)
        j.modalidade.setCurrentIndex(2)
        with patch('auditoria_fronteiras.fluxos.lembrar'):
            j.abrir_auditoria(self.base,self.uid)
        j.distancia_locais.setValue(0);j.extensao_locais.setValue(0)
        with patch('auditoria_fronteiras.interface.auditar',side_effect=AssertionError('núcleo proibido')),\
             patch.object(j,'salvar_operacao',return_value=False) as salvar:
            j.executar_auditoria()
        preparado=salvar.call_args.args[0]
        self.assertEqual(len(preparado['novo']['locais']),4)
        self.assertEqual(j.dossie['execucao_uid'],self.uid)
        j.close();projeto.clear()

    def test_reprocessar_formulario_chama_nucleo_uma_vez(self):
        projeto=QgsProject()
        for c in self.redes+self.limites:
            projeto.addMapLayer(c.camada)
        j=AuditoriaFronteirasDialog(projeto=projeto);j.modalidade.setCurrentIndex(3)
        with patch('auditoria_fronteiras.fluxos.lembrar'):
            j.abrir_auditoria(self.base,self.uid)
        j.tolerancia.setValue(.05)
        redes,limites,param=j.configuracoes()
        cfg=configuracao_proposta(redes,limites,param,j.opcoes_agrupamento())
        with abrir_ro(self.base) as con:
            esperado=__import__('auditoria_fronteiras.versoes',fromlist=['token']).token(con,self.uid)
        with patch.object(j,'preparar_reprocessamento',return_value=dict(configuracao=cfg,token_origem=esperado)),\
             patch('auditoria_fronteiras.interface.auditar',wraps=auditar) as nucleo,\
             patch.object(j,'salvar_operacao',return_value=False) as salvar,\
             patch('auditoria_fronteiras.interface.QMessageBox.critical') as erro:
            j.executar_auditoria()
        erro.assert_not_called();self.assertEqual(nucleo.call_count,1)
        self.assertEqual(salvar.call_args.args[0]['novo']['contexto']['parametros']['tolerancia_geometrica'],.05)
        self.assertEqual(j.dossie['execucao_uid'],self.uid)
        j.close();projeto.clear()

    def test_base_legada_migrada_apenas_na_gravacao(self):
        with conectar(self.base) as con:
            for t in ('AUDITORIA_PREFERENCIAS','AUDITORIA_LINHAGEM_AVALIADA','AUDITORIA_LINHAGEM_EVENTOS','AUDITORIA_LINHAGEM','AUDITORIA_VERSOES'):
                con.execute('DROP TABLE '+t)
                con.execute('DELETE FROM gpkg_contents WHERE table_name=?',(t,))
        antes=self.base.read_bytes()
        self.assertTrue(listar(self.base))
        self.assertEqual(self.base.read_bytes(),antes)
        s=self.gravar(self.preparar(),'SUBSTITUIR',confirmado=True)
        self.assertFalse(carregar_execucao(self.base,self.uid)['controle_versao']['ativa'])
        self.assertTrue(validar_base(s['backup']['backup']))

    def test_revisoes_afetadas_contam_historico_e_locais_separados(self):
        loc=self.revisar();anterior=historico(self.base,self.uid,loc)[0]['rev_uid']
        salvar_revisao(self.base,self.uid,loc,'EM_ANALISE','nova observação',esperada=anterior)
        c=self.preparar()['comparacao']
        self.assertEqual(c['revisoes_afetadas'],2);self.assertEqual(c['locais_revisados_afetados'],1)

    def test_campos_ausentes_exigem_remapeamento(self):
        projeto=QgsProject()
        for c in self.redes+self.limites:
            projeto.addMapLayer(c.camada)
        p=perfil(self.d);p['fontes'][0]['campo_id']='inexistente'
        with self.assertRaisesRegex(ValueError,'Campo ausente'):
            resolver_fontes(p,projeto)
        _,avisos=resolver_fontes(p,projeto,permitir_campos_ausentes=True)
        self.assertTrue(avisos);projeto.clear()

    def test_sessao_concorrente_substituiu_origem(self):
        p=self.preparar();q=self.preparar(1)
        self.gravar(q,'SUBSTITUIR',confirmado=True)
        with self.assertRaisesRegex(ValueError,'mudaram'):
            self.gravar(p,'SUBSTITUIR',confirmado=True)

    def test_fonte_realocada_preserva_chave_de_origem(self):
        from qgis.core import QgsVectorFileWriter
        from qgis.PyQt.QtCore import Qt
        projeto=QgsProject();p=perfil(self.d)
        for c in self.limites:
            projeto.addMapLayer(c.camada)
        arq=self.pasta/'rede_nova.gpkg'
        op=QgsVectorFileWriter.SaveVectorOptions();op.driverName='GPKG';op.layerName='rede'
        QgsVectorFileWriter.writeAsVectorFormatV3(self.redes[0].camada,str(arq),projeto.transformContext(),op)
        f=next(f for f in p['fontes'] if not f['modo']);f['fonte']='arquivo_ausente.shp';f['provedor']='ogr'
        fontes,_=resolver_fontes(p,projeto,lambda f:str(arq)+'|layername=rede')
        self.assertEqual(fontes[f['origem']].featureCount(),1)
        self.assertNotEqual(fontes[f['origem']].source(),f['fonte'])
        projeto.clear()

    def test_linhagem_externa_avaliacao_persistente(self):
        from auditoria_fronteiras.persistencia import avaliar_reconciliacoes
        p=self.preparar();outra=self.pasta/'outra.gpkg'
        novo=self.gravar(p,destino=outra)['dossie']
        self.assertTrue(novo['linhagem_locais'])
        l=novo['locais'][0]['uid']
        rev=salvar_revisao(outra,novo['execucao_uid'],l,'EM_ANALISE','avaliação da correspondência externa')
        avaliar_reconciliacoes(outra,novo['execucao_uid'],l,rev['uid'])
        reaberta=carregar_execucao(outra,novo['execucao_uid'])
        self.assertTrue(all(r['decisao']=='AVALIADA_MANUALMENTE' for r in reaberta['reconciliacoes'] if r['local_uid']==l))
        self.assertEqual(len(listar(self.base)),1)

    def test_primeira_persistencia_identifica_camadas_carregadas(self):
        projeto=QgsProject();j=AuditoriaFronteirasDialog(projeto=projeto)
        j.publicar_auditoria(self.d,None)
        self.assertFalse(carregadas(projeto)[0])
        with patch('auditoria_fronteiras.fluxos.lembrar'):
            j.vincular_base(self.base)
        self.assertEqual(carregadas(projeto)[0][0]['execucao_uid'],self.uid)
        self.assertEqual(j.base_revisao,str(self.base))
        j.close();projeto.clear()


if __name__=='__main__':
    unittest.main(verbosity=2)
