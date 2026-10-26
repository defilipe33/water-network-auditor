"""Inspeção, persistência, SHP/GPKG e outline XLSX no PyQGIS real."""
from copy import deepcopy
import json
from pathlib import Path
import sqlite3
import tempfile
import unittest
from unittest.mock import patch
import zipfile
import xml.etree.ElementTree as ET

from test_auditoria_fronteiras import APP,camada,quadrado
from qgis.core import QgsVectorLayer,QgsFeature,QgsGeometry,QgsProject,QgsVectorFileWriter
from qgis.PyQt.QtCore import Qt
from auditoria_fronteiras import auditar,CamadaEntrada,Parametros,Modo
from auditoria_fronteiras.investigacao import analisar_investigacao,OpcoesLocais
from auditoria_fronteiras.inspecao import construir_dossie,restaurar,detalhes_local,contagens
from auditoria_fronteiras.apresentacao import legivel
from auditoria_fronteiras.persistencia import (salvar_execucao,carregar_execucao,salvar_revisao,
    historico,listar_execucoes,conectar,exportar_gpkg)
from auditoria_fronteiras.persistencia import avaliar_reconciliacoes
from auditoria_fronteiras.exportacao import exportar_shp,validar_gpkg
from auditoria_fronteiras.painel import PainelInspecao,localizar_tubulacoes
from auditoria_fronteiras.visualizacao import criar_camadas_temporarias,carregar_no_projeto
from auditoria_fronteiras.relatorio import exportar_xlsx
from validar_estrutura_xlsx import validar_xlsx,S
from openpyxl import load_workbook


def rede(wkts,ids,nome='rede'):
    c=QgsVectorLayer('LineString?crs=EPSG:31983&field=id:string(1000)&field=sec:string&field=dn:double&field=mat:string',nome,'memory')
    fs=[]
    for w,id_ in zip(wkts,ids):
        f=QgsFeature(c.fields())
        f.setAttributes([id_,'Cadastro A',100.,'PVC'])
        f.setGeometry(QgsGeometry.fromWkt(w))
        fs.append(f)
    c.dataProvider().addFeatures(fs)
    return c


def preparar(duas=False,ids=None,distancia=5):
    linhas=['LINESTRING(-2 5,0 5)','LINESTRING(0 5,2 5)']
    ids=ids or ['00012','00013']
    redes=[rede(linhas,ids)] if not duas else [rede([linhas[0]],[ids[0]],'A'),rede([linhas[1]],[ids[1]],'B')]
    pol=camada('Polygon',[quadrado()],nome='limites')
    r=auditar([CamadaEntrada(c,'id') for c in redes],[CamadaEntrada(pol,'id','id',Modo.TRANSICAO)],Parametros())
    a=analisar_investigacao(r,OpcoesLocais(True,distancia,distancia))
    campos={f['origem']:dict(setor='sec',diametro='dn',material='mat') for f in r.fontes if not f['modo']}
    d=construir_dossie(r,a,campos)
    return r,a,d,redes,pol


class IdentificadoresTest(unittest.TestCase):
    def test_normalizacao_conservadora(self):
        for valor,esperado in [(1234567.0000000002,'1234567'),(12.,'12'),(12.25,'12.25'),
                (12.000000001,'12.000000001'),('0012','0012'),('12.0','12.0'),(None,'Sem ID'),
                (0,'0'),(9007199254740993,'9007199254740993'),(float('inf'),'inf')]:
            with self.subTest(valor=valor):
                self.assertEqual(legivel(valor),esperado)

    def test_correspondencia_mesma_camada_e_atributos(self):
        r,a,d,_,_=preparar()
        local=detalhes_local(d,d['locais'][0]['uid'])
        self.assertEqual(len(local['tubulacoes']),2)
        self.assertEqual({(t['id_exibido'],t['camada']) for t in local['tubulacoes']},
                         {('00012','rede'),('00013','rede')})
        self.assertEqual(len(d['evento_rede']),4)
        self.assertTrue(all(t['extras']['setor']=='Cadastro A' for t in local['tubulacoes']))
        self.assertEqual(a.locais[0]['registro']['REDE'],'rede; rede')

    def test_correspondencia_camadas_diferentes_ids_repetidos(self):
        _,_,d,_,_=preparar(True,ids=['001','001'])
        self.assertEqual(len({t['uid'] for t in d['tubulacoes']}),2)
        self.assertEqual({(t['id_exibido'],t['camada']) for t in d['tubulacoes']},{('001','A'),('001','B')})

    def test_setor_nao_configurado_nao_inferido(self):
        r,a,_,_,_=preparar()
        d=construir_dossie(r,a)
        self.assertTrue(all(not t['extras'] for t in d['tubulacoes']))
        self.assertTrue(any(t['regioes_geometricas'] for t in d['tubulacoes']))

    def test_bruto_preservado_e_snapshot_restaura_ids_e_geometria(self):
        r,a,_,_,_=preparar()
        r.ocorrencias[0].tubulacoes[0]['id_original']=1234567.0000000002
        a=analisar_investigacao(r,OpcoesLocais(True))
        d=construir_dossie(r,a)
        self.assertIn('1234567',[t['id_exibido'] for t in d['tubulacoes']])
        self.assertEqual(d['eventos'][0]['tecnico']['tubulacoes'][0]['id_original'],1234567.0000000002)
        rr,aa=restaurar(d)
        self.assertEqual([e.id for e in r.ocorrencias],[e.id for e in rr.ocorrencias])
        self.assertEqual([bytes(e.geometria.asWkb()) for e in r.ocorrencias],[bytes(e.geometria.asWkb()) for e in rr.ocorrencias])
        self.assertEqual(a.locais,aa.locais)

    def test_eventos_e_locais_ambiguos_separados(self):
        r,a,_,_,_=preparar()
        r.ocorrencias[0].ambigua=True
        r.ocorrencias[1].ambigua=True
        a=analisar_investigacao(r,OpcoesLocais(True))
        d=construir_dossie(r,a)
        self.assertEqual(contagens(d)['eventos_ambiguos'],2)
        self.assertEqual(contagens(d)['locais_ambiguos'],1)
        self.assertEqual(len(d['locais'][0]['eventos_ambiguos']),2)


class PersistenciaTest(unittest.TestCase):
    def setUp(self):
        # O pool OGR pode manter handles ociosos até o término do processo no Windows.
        self.tmp=tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        self.p=Path(self.tmp.name)/'revisao.gpkg'
        self.r,self.a,self.d,self.redes,self.pol=preparar()

    def tearDown(self):
        self.tmp.cleanup()

    def test_gpkg_relacional_reimportado(self):
        salvar_execucao(self.p,self.d)
        self.assertTrue(validar_gpkg(self.p,self.d)['valido'])
        with conectar(self.p) as con:
            self.assertEqual(con.execute('PRAGMA foreign_key_check').fetchall(),[])
            self.assertEqual(con.execute('SELECT count(*) FROM EVENTO_REDE').fetchone()[0],4)
            self.assertEqual(con.execute('SELECT count(*) FROM LOCAL_EVENTO').fetchone()[0],3)
            with self.assertRaises(sqlite3.IntegrityError):
                con.execute('INSERT INTO EVENTO_REDE VALUES(?,?,?)',(self.d['execucao_uid'],'inexistente','inexistente'))
            with self.assertRaises(sqlite3.IntegrityError):
                con.execute("INSERT INTO EVENTOS_FRONTEIRAS(EXE_UID,EVT_UID) VALUES('ausente','ausente')")

    def test_revisao_fecha_reabre_historico_e_concorrencia(self):
        salvar_execucao(self.p,self.d)
        exe=self.d['execucao_uid'];local=self.d['locais'][0]['uid']
        rev=salvar_revisao(self.p,exe,local,'EM_ANALISE','Verificar cadastro','Engenheiro')
        salvar_revisao(self.p,exe,local,'VERIFICADO','Inspeção documental','Engenheiro',rev['uid'])
        reaberta=carregar_execucao(self.p,exe)
        self.assertEqual(reaberta['locais'][0]['revisao']['situacao'],'VERIFICADO')
        self.assertEqual(len(historico(self.p,exe,local)),2)
        with self.assertRaises(ValueError):
            salvar_revisao(self.p,exe,local,'DESCARTADO',esperada=rev['uid'])
        with conectar(self.p) as con:
            self.assertEqual(con.execute('SELECT count(*) FROM REVISAO_EVENTO').fetchone()[0],6)
            with self.assertRaises(sqlite3.IntegrityError):
                con.execute('DELETE FROM REVISOES')

    def test_mesma_composicao_recupera_revisao_sem_apagar_historico(self):
        salvar_execucao(self.p,self.d)
        l=self.d['locais'][0]['uid']
        salvar_revisao(self.p,self.d['execucao_uid'],l,'VERIFICADO','Confirmado somente para este conjunto')
        novo=construir_dossie(self.r,self.a)
        recuperado=salvar_execucao(self.p,novo)
        self.assertEqual(recuperado['locais'][0]['revisao']['situacao'],'VERIFICADO')
        self.assertEqual(recuperado['reconciliacoes'][0]['tipo'],'EXATA')
        self.assertEqual(len(listar_execucoes(self.p)),2)
        self.assertEqual(len(historico(self.p,self.d['execucao_uid'],l)),1)

    def test_composicao_alterada_nao_herda_verificado(self):
        salvar_execucao(self.p,self.d)
        salvar_revisao(self.p,self.d['execucao_uid'],self.d['locais'][0]['uid'],'VERIFICADO')
        r,_=restaurar(self.d)
        r.ocorrencias=r.ocorrencias[:2]
        a=analisar_investigacao(r,OpcoesLocais(True))
        novo=construir_dossie(r,a)
        reaberto=salvar_execucao(self.p,novo)
        self.assertTrue(reaberto['reconciliacoes'])
        self.assertTrue(all(l['revisao']['situacao']=='PENDENTE' for l in reaberto['locais']))
        self.assertTrue(all(x['tipo']=='PARCIAL' for x in reaberto['reconciliacoes']))

    def test_backup_preserva_historico_integral(self):
        salvar_execucao(self.p,self.d)
        l=self.d['locais'][0]['uid'];exe=self.d['execucao_uid']
        rev=salvar_revisao(self.p,exe,l,'EM_ANALISE','1')
        salvar_revisao(self.p,exe,l,'VERIFICADO','2',esperada=rev['uid'])
        copia=Path(self.tmp.name)/'entrega.gpkg'
        exportar_gpkg(copia,carregar_execucao(self.p,exe),self.p)
        self.assertEqual(len(historico(copia,exe,l)),2)
        self.assertTrue(validar_gpkg(copia,carregar_execucao(self.p,exe))['valido'])

    def test_idempotencia_e_base_alheia_preservada(self):
        salvar_execucao(self.p,self.d)
        salvar_execucao(self.p,self.d)
        self.assertEqual(len(listar_execucoes(self.p)),1)
        alheia=Path(self.tmp.name)/'alheia.gpkg'
        alheia.write_bytes(b'nao alterar')
        with self.assertRaises(ValueError):
            salvar_execucao(alheia,self.d)
        self.assertEqual(alheia.read_bytes(),b'nao alterar')

    def test_situacao_invalida_e_local_ausente(self):
        salvar_execucao(self.p,self.d)
        with self.assertRaises(ValueError):
            salvar_revisao(self.p,self.d['execucao_uid'],self.d['locais'][0]['uid'],'ABERTA')
        with self.assertRaises(sqlite3.IntegrityError):
            salvar_revisao(self.p,self.d['execucao_uid'],'ausente','PENDENTE')

    def test_reabertura_colunas_status_e_exportacao_sem_depender_ordem_json(self):
        salvar_execucao(self.p,self.d)
        uid=self.d['locais'][0]['uid'];exe=self.d['execucao_uid']
        salvar_revisao(self.p,exe,uid,'VERIFICADO','Revisão de teste')
        d=carregar_execucao(self.p,exe)
        r,a=restaurar(d)
        camadas=criar_camadas_temporarias(r,analise=a)
        local=next(camadas[2].getFeatures())
        self.assertEqual(local['N_EVENTOS'],3)
        self.assertEqual(local['LOCAL'],self.d['locais'][0]['registro']['LOCAL'])
        p=Path(self.tmp.name)/'reaberto.xlsx'
        exportar_xlsx(r,p,analise=a,dossie=d)
        wb=load_workbook(p)
        self.assertEqual(wb['PONTOS_INVESTIGACAO']['B2'].value,3)
        self.assertEqual(wb['PONTOS_INVESTIGACAO']['N2'].value,'VERIFICADO')
        self.assertEqual(wb['PONTOS_INVESTIGACAO']['L2'].value,0)
        wb.close()

    def test_reconciliacao_explicita_e_auditavel(self):
        salvar_execucao(self.p,self.d)
        salvar_revisao(self.p,self.d['execucao_uid'],self.d['locais'][0]['uid'],'VERIFICADO','Anterior')
        r,_=restaurar(self.d)
        r.ocorrencias=r.ocorrencias[:2]
        d=construir_dossie(r,analisar_investigacao(r,OpcoesLocais(True)))
        novo=salvar_execucao(self.p,d)
        local=novo['locais'][0]['uid'];exe=novo['execucao_uid']
        rev=salvar_revisao(self.p,exe,local,'EM_ANALISE','Composição alterada examinada')
        avaliar_reconciliacoes(self.p,exe,local,rev['uid'])
        self.assertEqual(carregar_execucao(self.p,exe)['reconciliacoes'][0]['decisao'],'AVALIADA_MANUALMENTE')
        self.assertEqual(historico(self.p,self.d['execucao_uid'],self.d['locais'][0]['uid'])[0]['situacao'],'VERIFICADO')

    def test_crs_divergente_nao_mistura_execucoes(self):
        salvar_execucao(self.p,self.d)
        r,_=restaurar(self.d)
        from dataclasses import replace
        r.parametros=replace(r.parametros,crs='EPSG:31982')
        d=construir_dossie(r,analisar_investigacao(r,OpcoesLocais(True)))
        with self.assertRaises(ValueError):
            salvar_execucao(self.p,d)
        self.assertEqual(len(listar_execucoes(self.p)),1)

    def test_geometria_corrigida_e_candidata_sem_herdar_decisao(self):
        salvar_execucao(self.p,self.d)
        salvar_revisao(self.p,self.d['execucao_uid'],self.d['locais'][0]['uid'],'VERIFICADO')
        c=self.redes[0]
        novas={f.id():QgsGeometry.fromWkt('LINESTRING(-2 5.001,0 5.001)' if n==0 else
              'LINESTRING(0 5.001,2 5.001)') for n,f in enumerate(c.getFeatures())}
        c.dataProvider().changeGeometryValues(novas)
        r=auditar([CamadaEntrada(c,'id')],[CamadaEntrada(self.pol,'id','id',Modo.TRANSICAO)])
        a=analisar_investigacao(r,OpcoesLocais(True))
        d=construir_dossie(r,a)
        self.assertFalse({e['uid'] for e in self.d['eventos']} & {e['uid'] for e in d['eventos']})
        atual=salvar_execucao(self.p,d)
        self.assertTrue(atual['reconciliacoes'])
        self.assertTrue(all(r['tipo']=='PARCIAL' for r in atual['reconciliacoes']))
        self.assertEqual(atual['locais'][0]['revisao']['situacao'],'PENDENTE')


class ExportacaoTest(unittest.TestCase):
    def test_exportacao_sem_agrupamento_inclui_camadas_vazias_validas(self):
        r,_,_,_,_=preparar()
        a=analisar_investigacao(r,OpcoesLocais(False))
        d=construir_dossie(r,a)
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as pasta:
            p=Path(pasta)
            salvar_execucao(p/'sem_locais.gpkg',d)
            self.assertTrue(validar_gpkg(p/'sem_locais.gpkg',d)['valido'])
            s=exportar_shp(d,p/'SHP')
            self.assertEqual(s['validacoes'][1]['registros'],0)
            exportar_xlsx(r,p/'sem_locais.xlsx',analise=a,dossie=d)
            self.assertTrue(validar_xlsx(p/'sem_locais.xlsx')['valido'])

    def test_coincidencia_preserva_wkb_linear_e_exporta_ponto_representativo(self):
        c=rede(['LINESTRING(0 2,0 8)'],['001'])
        pol=camada('Polygon',[quadrado()])
        r=auditar([CamadaEntrada(c,'id')],[CamadaEntrada(pol,'id','id')])
        a=analisar_investigacao(r,OpcoesLocais(True))
        d=construir_dossie(r,a)
        rr,_=restaurar(d)
        self.assertTrue(any(e.classificacao=='COINCIDENCIA' for e in rr.ocorrencias))
        self.assertEqual([bytes(e.geometria.asWkb()) for e in r.ocorrencias],
                         [bytes(e.geometria.asWkb()) for e in rr.ocorrencias])
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as pasta:
            p=Path(pasta)
            salvar_execucao(p/'coincidencia.gpkg',d)
            self.assertTrue(validar_gpkg(p/'coincidencia.gpkg',d)['valido'])
            exportar_shp(d,p/'SHP')

    def test_shp_geometria_crs_campos_auxiliares_e_zeros(self):
        r,a,d,_,_=preparar(True,ids=['00001','00001'])
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as pasta:
            shp=exportar_shp(d,Path(pasta)/'SHP')
            self.assertEqual(shp['avisos'],[])
            self.assertEqual(shp['validacoes'][0]['registros'],3)
            camada=QgsVectorLayer(str(Path(pasta)/'SHP/EVENTOS_FRONTEIRAS.shp'),'e','ogr')
            self.assertTrue(all(len(n)<=10 for n in camada.fields().names()))
            self.assertTrue(all('00001' in f['ID_REDE'] for f in camada.getFeatures()))
            from qgis.PyQt import sip
            sip.delete(camada)
            gpkg=Path(pasta)/'base.gpkg'
            salvar_execucao(gpkg,d)
            self.assertTrue(validar_gpkg(gpkg,d)['valido'])
            self.assertTrue((Path(pasta)/'SHP/EVENTO_REDE.csv').exists())
            self.assertTrue((Path(pasta)/'SHP/LOCAL_EVENTO.csv').exists())
            self.assertTrue((Path(pasta)/'SHP/EVENTO_POLIGONO.csv').exists())

    def test_shp_valores_longos_nulos_grandes_sem_truncamento(self):
        for ids in ([None,'9007199254740993001'],['é'*200,'0000']):
            with self.subTest(ids=ids),tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as pasta:
                _,_,d,_,_=preparar(ids=ids)
                x=exportar_shp(d,Path(pasta)/'SHP')
                if ids[0]:
                    self.assertTrue(x['avisos'])
                bruto=json.loads((Path(pasta)/'SHP/DOSSIE_TECNICO.json').read_text(encoding='utf-8'))
                self.assertEqual([t['id_bruto'] for t in bruto['tubulacoes']],ids)

    def test_nao_sobrescreve_pasta_existente(self):
        _,_,d,_,_=preparar()
        with tempfile.TemporaryDirectory() as p:
            with self.assertRaises(FileExistsError):
                exportar_shp(d,p)

    def test_outline_xml_dados_resumo_e_revisao(self):
        r,a,d,_,_=preparar(True)
        l=d['locais'][0]
        l['revisao']['situacao']='EM_ANALISE'
        l['registro']['SITUACAO']='EM_ANALISE'
        with tempfile.TemporaryDirectory() as pasta:
            saida=exportar_xlsx(r,Path(pasta)/'expansivel.xlsx',analise=a,dossie=d)
            self.assertTrue(validar_xlsx(saida['xlsx'])['valido'])
            wb=load_workbook(saida['xlsx'])
            ws=wb['DETALHAMENTO_LOCAIS']
            self.assertFalse(ws.tables)
            self.assertIsNone(ws.auto_filter.ref)
            self.assertFalse(ws.sheet_properties.outlinePr.summaryBelow)
            self.assertTrue(ws.row_dimensions[2].collapsed)
            self.assertTrue(ws.row_dimensions[3].hidden)
            self.assertEqual(ws.row_dimensions[3].outlineLevel,1)
            rows=list(ws.values)
            self.assertEqual(sum(r[0]=='TUBULACAO' for r in rows),2)
            self.assertEqual(sum(r[0]=='EVENTO' for r in rows),3)
            self.assertEqual(sum(r[0]=='AREA' for r in rows),1)
            self.assertEqual(rows[1][9],'EM_ANALISE')
            self.assertEqual(wb['PONTOS_INVESTIGACAO']['N2'].value,'EM_ANALISE')
            wb.close()
            with zipfile.ZipFile(saida['xlsx']) as z:
                xml=ET.fromstring(z.read('xl/worksheets/sheet8.xml'))
                self.assertEqual(xml.find(f'{{{S}}}sheetPr/{{{S}}}outlinePr').get('summaryBelow'),'0')
                self.assertTrue(xml.findall(f'{{{S}}}sheetData/{{{S}}}row[@outlineLevel="1"]'))


class PainelTest(unittest.TestCase):
    def setUp(self):
        self.alertas=patch('auditoria_fronteiras.painel.QMessageBox.warning').start()
        from qgis.PyQt.QtWidgets import QMessageBox
        self.perguntas=patch('auditoria_fronteiras.painel.QMessageBox.question',return_value=QMessageBox.Discard).start()
        self.r,self.a,self.d,self.redes,self.pol=preparar(True)
        self.projeto=QgsProject()
        for c in self.redes+[self.pol]:
            self.projeto.addMapLayer(c)
        self.camadas=criar_camadas_temporarias(self.r,analise=self.a)
        carregar_no_projeto(self.camadas,self.projeto)
        self.painel=PainelInspecao(self.d,self.camadas,self.projeto)

    def tearDown(self):
        self.painel.close()
        self.painel.deleteLater()
        APP.processEvents()
        self.projeto.clear()
        patch.stopall()

    def test_secoes_expansiveis_e_campos(self):
        self.painel.locais.setCurrentIndex(1)
        arv=self.painel.arvore
        self.assertEqual(arv.topLevelItemCount(),4)
        self.assertEqual(arv.topLevelItem(0).childCount(),2)
        self.assertEqual(arv.topLevelItem(1).childCount(),3)
        self.assertEqual(arv.topLevelItem(2).childCount(),1)
        item=arv.topLevelItem(0)
        self.assertFalse(item.isExpanded())
        item.setExpanded(True)
        self.assertTrue(item.isExpanded())
        self.assertIn('Cadastro A',item.child(0).text(3))

    def test_selecao_local_eventos_e_inversa(self):
        self.painel.locais.setCurrentIndex(1)
        self.assertEqual(self.camadas[0].selectedFeatureCount(),3)
        self.assertEqual(self.camadas[2].selectedFeatureCount(),1)
        self.painel.locais.setCurrentIndex(0)
        self.camadas[2].removeSelection()
        self.camadas[2].selectAll()
        self.assertEqual(self.painel.locais.currentIndex(),1)
        item=self.painel.arvore.topLevelItem(1).child(0)
        self.painel.arvore.setCurrentItem(item)
        self.assertEqual(self.camadas[0].selectedFeatureCount(),1)

    def test_origens_por_assinatura_nao_id_isolado(self):
        self.painel.locais.setCurrentIndex(1)
        self.painel.selecionar_origens()
        self.assertEqual([c.selectedFeatureCount() for c in self.redes],[1,1])

    def test_ids_repetidos_na_mesma_origem_encontram_geometrias_distintas(self):
        _,_,d,redes,pol=preparar(ids=['001','001'])
        p=QgsProject()
        p.addMapLayer(redes[0])
        encontrados,avisos=localizar_tubulacoes(d,d['tubulacoes'],p)
        self.assertFalse(avisos)
        self.assertEqual(len(encontrados[redes[0].id()]),2)
        p.clear()

    def test_copias_identicas_nao_escolhidas_arbitrariamente(self):
        c=rede(['LINESTRING(-2 5,2 5)']*2,['001']*2)
        p=QgsProject()
        p.addMapLayer(c)
        pol=camada('Polygon',[quadrado()])
        r=auditar([CamadaEntrada(c,'id')],[CamadaEntrada(pol,'id','id')])
        a=analisar_investigacao(r,OpcoesLocais(True))
        d=construir_dossie(r,a)
        encontrados,avisos=localizar_tubulacoes(d,d['tubulacoes'],p)
        self.assertTrue(avisos)
        self.assertEqual(encontrados[c.id()],[])
        p.clear()

    def test_origens_ausentes_nao_impedem_consulta(self):
        for c in self.redes:
            self.projeto.removeMapLayer(c.id())
        self.painel.locais.setCurrentIndex(1)
        self.painel.selecionar_origens()
        self.assertIn('ausente',self.painel.aviso.text())
        self.assertEqual(self.painel.arvore.topLevelItem(0).childCount(),2)

    def test_resultados_removidos_nao_impedem_consulta_preservada(self):
        for c in self.camadas:
            self.projeto.removeMapLayer(c.id())
        self.painel.locais.setCurrentIndex(1)
        self.painel.destacar()
        self.assertEqual(self.painel.arvore.topLevelItem(1).childCount(),3)

    def test_edicoes_nao_salvas_podem_cancelar_troca_de_local(self):
        from qgis.PyQt.QtWidgets import QMessageBox
        self.painel.locais.setCurrentIndex(1)
        self.painel.observacoes.setPlainText('Ainda não salvo')
        with patch('auditoria_fronteiras.painel.QMessageBox.question',return_value=QMessageBox.Cancel):
            self.painel.locais.setCurrentIndex(0)
        self.assertEqual(self.painel.locais.currentIndex(),1)
        self.assertEqual(self.painel.observacoes.toPlainText(),'Ainda não salvo')
        with patch('auditoria_fronteiras.painel.QMessageBox.question',return_value=QMessageBox.Discard):
            self.painel.locais.setCurrentIndex(0)
        self.assertFalse(self.painel.tem_edicoes())

    def test_salvar_revisao_e_consultar_nova_sessao(self):
        with tempfile.TemporaryDirectory() as pasta:
            p=Path(pasta)/'rev.gpkg'
            self.painel.locais.setCurrentIndex(1)
            self.painel.situacao.setCurrentText('EM_ANALISE')
            self.painel.observacoes.setPlainText('Inspecionar as duas redes')
            with patch('auditoria_fronteiras.painel.QFileDialog.getSaveFileName',return_value=(str(p),'')):
                self.painel.gravar_revisao()
            l=carregar_execucao(p,self.d['execucao_uid'])['locais'][0]
            self.assertEqual(l['revisao']['situacao'],'EM_ANALISE')
            self.assertEqual(l['revisao']['observacoes'],'Inspecionar as duas redes')
            self.assertEqual(next(self.camadas[2].getFeatures())['SITUACAO'],'EM_ANALISE')


if __name__=='__main__':
    unittest.main(verbosity=2)
