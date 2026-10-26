"""Agrupamento conservador, vínculos e relatórios; sem dados reais fixados."""
import json
from pathlib import Path
import tempfile
import unittest
from copy import deepcopy
from unittest.mock import patch

from test_auditoria_fronteiras import APP,camada,quadrado
from qgis.core import QgsGeometry,QgsProject
from qgis.PyQt.QtCore import Qt
from auditoria_fronteiras.modelos import Ocorrencia,Resultado,Parametros,Modo,Cancelado
from auditoria_fronteiras.investigacao import (OpcoesLocais,analisar_investigacao,
                                             validar_analise)
from auditoria_fronteiras.visualizacao import criar_camadas_temporarias,carregar_no_projeto
from auditoria_fronteiras.interface import AuditoriaFronteirasDialog
from auditoria_fronteiras.relatorio import exportar_xlsx
from auditoria_fronteiras import CamadaEntrada,auditar
from validar_estrutura_xlsx import validar_xlsx
from validar_relatorio_locais import validar as validar_relatorio_locais
from openpyxl import load_workbook


def ref(chave,origem='lim',nome=None):
    return dict(chave=chave,origem=origem,camada=origem,id_original=chave,
                nome=nome or chave,atributos_originais={})


def evento(uid,x=0,y=0,tubos=('T1',),pol='P1',origem='lim',classe='TRAVESSIA'):
    return Ocorrencia(uid,classe,Modo.PERIMETRO,QgsGeometry.fromWkt(f'POINT({x} {y})'),
        [ref(t,'rede') for t in tubos],[ref(pol,origem)],0,x)


def resultado(eventos):
    r=Resultado(Parametros(),ocorrencias=eventos)
    origens=sorted({p['origem'] for e in eventos for p in e.poligonos})
    r.fontes=[dict(origem=o,camada=o,fonte='memory:'+o,modo='PERIMETRO',
                   campo_id='id',campo_nome='nome',crs_origem='EPSG:31983') for o in origens]
    r.contadores={f'2:{o}':dict(lidas=2,ids_repetidos=0) for o in origens}
    return r


class LocaisTest(unittest.TestCase):
    def agrupar(self,r,distancia=5,extensao=5):
        return analisar_investigacao(r,OpcoesLocais(True,distancia,extensao))

    def test_coincidentes_mesma_fronteira_mesma_classe(self):
        a=self.agrupar(resultado([evento('a'),evento('b',tubos=('T2',))]))
        self.assertEqual(len(a.locais),1)
        self.assertEqual(a.locais[0]['registro']['N_EVENTOS'],2)
        self.assertEqual(a.locais[0]['registro']['N_TUBOS'],2)

    def test_proximos_compativeis_e_incompativeis(self):
        self.assertEqual(len(self.agrupar(resultado([evento('a'),evento('b',2)])).locais),1)
        self.assertEqual(len(self.agrupar(resultado([evento('a'),evento('b',2,tubos=('T2',))])).locais),2)

    def test_distancia_sozinha_nao_basta_nem_com_coordenadas_iguais(self):
        r=resultado([evento('a'),evento('b',tubos=('T2',),pol='P2',origem='outra')])
        self.assertEqual(len(self.agrupar(r).locais),2)

    def test_duas_extremidades_e_testemunho_de_continuidade(self):
        a=evento('a',-.2,classe='EXTREMIDADE',tubos=('A',))
        b=evento('b',.2,classe='EXTREMIDADE',tubos=('B',))
        c=evento('c',classe='CONTINUIDADE_POTENCIAL',tubos=('A','B'))
        c.detalhes={'extremidades':[dict(tubulacao='A',parte=0,lado='FIM',x=-.2,y=0),
                                   dict(tubulacao='B',parte=0,lado='INICIO',x=.2,y=0)]}
        i=self.agrupar(resultado([a,b,c]),1,1)
        self.assertEqual(len(i.locais),1)
        self.assertEqual(i.locais[0]['registro']['N_EVENTOS'],3)
        self.assertEqual(i.locais[0]['registro']['N_TUBOS'],2)
        self.assertEqual(len(i.vinculos),3)

    def test_analises_distintas_mantem_todos_poligonos(self):
        r=resultado([evento('a'),evento('b',.5,pol='P2',origem='outra')])
        i=self.agrupar(r)
        self.assertEqual(len(i.locais),1)
        self.assertEqual(set(i.locais[0]['poligonos']),{'P1','P2'})

    def test_aninhados_reais_e_preservacao_do_nucleo(self):
        rede=camada('LineString',['LINESTRING(-2 5,12 5)'])
        a=camada('Polygon',[quadrado()])
        b=camada('Polygon',[quadrado(2,2,6)])
        r=auditar([CamadaEntrada(rede,'id')],[CamadaEntrada(a,'id'),CamadaEntrada(b,'id')])
        antes=[(e.id,e.classificacao,bytes(e.geometria.asWkb()),deepcopy(e.poligonos)) for e in r.ocorrencias]
        i=self.agrupar(r,3,3)
        self.assertEqual(len(r.ocorrencias),4)
        self.assertEqual(len(i.locais),2)
        self.assertEqual(antes,[(e.id,e.classificacao,bytes(e.geometria.asWkb()),e.poligonos) for e in r.ocorrencias])

    def test_poligonos_sobrepostos_preservam_candidatos(self):
        rede=camada('LineString',['LINESTRING(-2 5,17 5)'])
        p=camada('Polygon',[quadrado(),quadrado(5)])
        r=auditar([CamadaEntrada(rede,'id')],[CamadaEntrada(p,'id','nome',Modo.TRANSICAO)])
        candidatos=[deepcopy(e.regiao_depois) for e in r.ocorrencias]
        i=self.agrupar(r)
        self.assertEqual(len(i.vinculos),4)
        self.assertEqual(candidatos,[e.regiao_depois for e in r.ocorrencias])
        self.assertTrue(any(p['LEITURA']=='MULTIPLA' for p in i.pares))

    def test_cadeia_nao_e_componente_transitiva(self):
        i=self.agrupar(resultado([evento('a',0),evento('b',4),evento('c',8)]))
        self.assertEqual(len(i.locais),2)
        self.assertTrue(all(l['extensao_m']<=5 for l in i.locais))

    def test_extensao_maxima_independente_do_raio_da_semente(self):
        r=resultado([evento('a',4),evento('b',0),evento('c',8)])
        self.assertEqual(len(self.agrupar(r,5,5).locais),2)
        self.assertEqual(len(self.agrupar(r,5,10).locais),1)

    def test_ids_estaveis_independentes_da_ordem_de_entrada(self):
        r=resultado([evento('a'),evento('b',2),evento('c',12)])
        a=self.agrupar(r)
        r.ocorrencias.reverse()
        b=self.agrupar(r)
        self.assertEqual({l['uid']:l['eventos'] for l in a.locais},{l['uid']:l['eventos'] for l in b.locais})
        self.assertEqual({v['evento_uid']:v['local_uid'] for v in a.vinculos},
                         {v['evento_uid']:v['local_uid'] for v in b.vinculos})

    def test_vinculos_integros_e_adulteracao_rejeitada(self):
        r=resultado([evento('a'),evento('b',10)])
        i=self.agrupar(r)
        validar_analise(r,i)
        i.vinculos[0]['local_uid']=i.vinculos[1]['local_uid']
        with self.assertRaises(ValueError):
            validar_analise(r,i)

    def test_classificacoes_incompativeis_e_trecho_extenso(self):
        r=resultado([evento('a'),evento('b',classe='TANGENCIA')])
        self.assertEqual(len(self.agrupar(r).locais),2)
        a,b=evento('a',classe='COINCIDENCIA'),evento('b',classe='COINCIDENCIA')
        a.geometria=QgsGeometry.fromWkt('LINESTRING(0 0,10 0)')
        b.geometria=QgsGeometry(a.geometria)
        self.assertEqual(len(self.agrupar(resultado([a,b])).locais),2)

    def test_ids_nulos_ou_iguais_nao_colapsam_tubulacoes(self):
        a,b=evento('a'),evento('b',tubos=('T2',))
        a.tubulacoes[0]['id_original']=None
        b.tubulacoes[0]['id_original']=None
        i=self.agrupar(resultado([a,b]))
        self.assertEqual(i.locais[0]['registro']['N_TUBOS'],2)
        self.assertEqual(i.locais[0]['registro']['ID_REDE'],'Sem ID; Sem ID')

    def test_medóide_e_um_evento_e_nao_move_os_demais(self):
        r=resultado([evento('a',0),evento('b',1),evento('c',4)])
        i=self.agrupar(r)
        self.assertEqual(i.locais[0]['representante'],'b')
        self.assertEqual(i.locais[0]['registro']['X'],1)
        self.assertEqual([e.geometria.asPoint().x() for e in r.ocorrencias],[0,1,4])

    def test_sem_agrupamento_e_cancelamento(self):
        r=resultado([evento('a')])
        i=analisar_investigacao(r)
        self.assertEqual(i.locais,[])
        self.assertEqual(i.vinculos,[])
        with self.assertRaises(Cancelado):
            analisar_investigacao(r,OpcoesLocais(True),cancelar=lambda:True)
        with self.assertRaises(ValueError):
            self.agrupar(r,-1,5)

    def test_camadas_e_relacoes_qgis(self):
        r=resultado([evento('a'),evento('b',2)])
        i=self.agrupar(r)
        camadas=criar_camadas_temporarias(r,analise=i)
        self.assertEqual([c.name() for c in camadas],
            ['EVENTOS_FRONTEIRAS','DETALHES_TECNICOS','PONTOS_INVESTIGACAO','VINCULOS_LOCAIS'])
        self.assertEqual(camadas[0].featureCount(),2)
        self.assertEqual(camadas[2].featureCount(),1)
        self.assertEqual(camadas[3].featureCount(),2)
        self.assertEqual(next(camadas[2].getFeatures())['SITUACAO'],'PENDENTE')
        projeto=QgsProject()
        carregar_no_projeto(camadas,projeto)
        self.assertEqual(len(projeto.relationManager().relations()),3)
        self.assertTrue(all(rel.isValid() for rel in projeto.relationManager().relations().values()))
        projeto.clear()


class SintesesRelatorioTest(unittest.TestCase):
    def transicao(self,uid,antes,depois,classe='TRAVESSIA'):
        e=evento(uid,classe=classe)
        e.modo=Modo.TRANSICAO
        e.regiao_antes=antes
        e.regiao_depois=depois
        return e

    def test_pares_independentes_da_orientacao(self):
        a,b=ref('A'),ref('B')
        r=resultado([self.transicao('1',[a],[b]),self.transicao('2',[b],[a])])
        i=analisar_investigacao(r)
        self.assertEqual(len(i.pares),1)
        self.assertEqual(i.pares[0]['N_EVENTOS'],2)
        self.assertEqual(set(i.tecnicos_pares[0]['eventos']),{'1','2'})

    def test_sobreposicao_sem_atribuicao_arbitraria(self):
        a,b,c=ref('A'),ref('B'),ref('C')
        i=analisar_investigacao(resultado([self.transicao('1',[a,b],[c])]))
        self.assertEqual(i.contadores['pares_inequivocos'],0)
        self.assertEqual(i.pares[0]['LEITURA'],'MULTIPLA')
        self.assertEqual(i.pares[0]['REGIAO_A'],'A; B')
        self.assertEqual(len(i.tecnicos_pares[0]['antes']),2)

    def test_exterior_nao_e_extremidade_ou_indeterminado(self):
        a=ref('A')
        verdadeiro=self.transicao('1',[a],[])
        extremidade=self.transicao('2',[a],[],classe='EXTREMIDADE')
        extremidade.lado_depois_disponivel=False
        ambiguo=self.transicao('3',[a],[],classe='AMBIGUA')
        ambiguo.ambigua=True
        i=analisar_investigacao(resultado([verdadeiro,extremidade,ambiguo]))
        self.assertEqual(len(i.saidas),1)
        self.assertEqual(i.saidas[0]['uid'],'1')
        self.assertEqual(i.pares[0]['LEITURA'],'INDETERMINADA')

    def test_metadados_variacoes_de_escopo_e_parametros(self):
        a=resultado([evento('a')])
        b=resultado([evento('a'),evento('b',origem='nova')])
        ia=analisar_investigacao(a,OpcoesLocais(True,3,4),data_execucao='2026-10-24T12:00:00+00:00')
        ib=analisar_investigacao(b,OpcoesLocais(False,6,7))
        self.assertEqual(ia.metadados['data_execucao_utc'],'2026-10-24T12:00:00+00:00')
        self.assertEqual(ib.metadados['data_execucao_utc'],'Não registrada')
        self.assertEqual(len(ib.metadados['fontes']),2)
        self.assertNotEqual(ia.assinatura_eventos,ib.assinatura_eventos)
        self.assertEqual(ia.metadados['agrupamento']['distancia'],3)
        with self.assertRaises(ValueError):
            validar_analise(b,ia)

    def test_excel_contagens_dados_vinculos_e_xml(self):
        r=resultado([evento('a'),evento('b',2),evento('c',12)])
        i=analisar_investigacao(r,OpcoesLocais(True),data_execucao='2026-10-24T12:00:00+00:00')
        with tempfile.TemporaryDirectory() as pasta:
            arq=exportar_xlsx(r,Path(pasta)/'locais.xlsx',analise=i)
            self.assertTrue(validar_xlsx(arq['xlsx'])['valido'])
            self.assertTrue(validar_relatorio_locais(arq['xlsx'])['vinculos_integros'])
            wb=load_workbook(arq['xlsx'])
            self.assertEqual(wb['PONTOS_INVESTIGACAO'].max_row-1,2)
            self.assertEqual(sum(row[1] for row in wb['PONTOS_INVESTIGACAO'].iter_rows(min_row=2,values_only=True)),3)
            self.assertTrue(any(row[3]=='2026-10-24T12:00:00+00:00'
                for row in wb['PARAMETROS'].iter_rows(min_row=2,values_only=True)))
            wb.close()
            dados=json.loads(Path(arq['tecnico']).read_text(encoding='utf-8'))
            self.assertEqual(len(dados['ocorrencias']),3)
            self.assertEqual(len(dados['investigacao']['vinculos']),3)

    def test_interface_opcional_e_metadados_na_execucao(self):
        for metodo in ('confirmar_processamento','resolver_pendencias'):
            controle=patch.object(AuditoriaFronteirasDialog,metodo,return_value=True)
            controle.start();self.addCleanup(controle.stop)
        p=QgsProject()
        r=camada('LineString',['LINESTRING(-2 5,12 5)'])
        l=camada('Polygon',[quadrado()])
        p.addMapLayer(r)
        p.addMapLayer(l)
        janela=AuditoriaFronteirasDialog(projeto=p)
        for tabela in (janela.redes,janela.limites):
            tabela.item(0,0).setCheckState(Qt.Checked)
        janela.agrupar.setChecked(True)
        janela.distancia_locais.setValue(5)
        janela.extensao_locais.setValue(5)
        with patch('auditoria_fronteiras.interface.QMessageBox.critical') as falha:
            janela.executar_auditoria()
        falha.assert_not_called()
        self.assertEqual(len(janela.camadas_geradas),4)
        self.assertNotEqual(janela.analise.metadados['data_execucao_utc'],'Não registrada')
        self.assertIn('Locais de investigação',janela.resumo.toPlainText())
        janela.iniciar_operacao(0)
        for tabela in (janela.redes,janela.limites):
            tabela.item(0,0).setCheckState(Qt.Checked)
        janela.agrupar.setChecked(False)
        with patch('auditoria_fronteiras.interface.QMessageBox.critical'):
            janela.executar_auditoria()
        self.assertEqual(len(janela.camadas_geradas),2)
        janela.close()
        p.clear()


if __name__=='__main__':
    unittest.main(verbosity=2)
