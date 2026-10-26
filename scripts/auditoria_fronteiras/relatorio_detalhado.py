"""Outline nativo Excel: um nível de detalhes, resumos acima, sem tabelas/filtros."""
from .inspecao import detalhes_local
from .apresentacao import legivel,vazio

CAMPOS=('TIPO','LOCAL','ID','CAMADA','CLASSIF','AREA','AMBIGUO','N_EVENTOS','N_TUBOS',
        'SITUACAO','SETOR_CAD','DIAMETRO','MATERIAL','INFORMACAO','TEC_UID')


def criar_detalhamento(wb,dossie,planilha):
    linhas=[]
    grupos=[]
    fontes={f['origem']:f for f in dossie['contexto']['fontes']}
    def adicionar(**dados):
        linhas.append([dados.get(c,'') for c in CAMPOS])
    for l in dossie['locais']:
        d=detalhes_local(dossie,l['uid'])
        r=l['registro']
        principal=len(linhas)+2
        adicionar(TIPO='LOCAL',LOCAL=r['LOCAL'],CLASSIF=r['CLASSIF'],AREA=r['AREAS'],AMBIGUO=r['AMBIGUO'],
            N_EVENTOS=r['N_EVENTOS'],N_TUBOS=r['N_TUBOS'],SITUACAO=l['revisao']['situacao'],TEC_UID=l['uid'],
            INFORMACAO=l['revisao']['observacoes'])
        for t in d['tubulacoes']:
            extra=t['extras']
            adicionar(TIPO='TUBULACAO',LOCAL=r['LOCAL'],ID=t['id_exibido'],CAMADA=t['camada'],
                SETOR_CAD=legivel(extra.get('setor'),''),DIAMETRO=legivel(extra.get('diametro'),''),
                MATERIAL=legivel(extra.get('material'),''),TEC_UID=t['uid'],
                INFORMACAO='Regiões geométricas observadas nas fronteiras: '+ '; '.join(
                    legivel(p.get('nome') or p.get('id_original')) for p in t['regioes_geometricas']))
        for e in d['eventos']:
            er=e['registro']
            adicionar(TIPO='EVENTO',LOCAL=r['LOCAL'],ID=er['ID'],CAMADA=er['TIPO_LIM'],CLASSIF=er['CLASSIF'],
                AREA=er['AREA'],AMBIGUO=er['AMBIGUO'],TEC_UID=e['uid'],
                INFORMACAO=f"Tubulações: {er['ID_REDE']} | Antes: {er['REG_ANTES']} | Depois: {er['REG_DEPOIS']}")
        for p in d['poligonos']:
            fonte=fontes.get(p['origem'],{})
            exibicao=p.get('nome') if fonte.get('campo_nome') else p.get('id_original')
            if vazio(exibicao):
                exibicao=p.get('id_original')
            associados=[e for e in d['eventos'] if any(x['evento_uid']==e['uid'] and x['pol_uid']==p['chave']
                        for x in dossie['evento_poligono'])]
            trans='; '.join(dict.fromkeys(f"{e['registro']['REG_ANTES']} → {e['registro']['REG_DEPOIS']}"
                for e in associados if e['tecnico']['modo']=='TRANSICAO'))
            adicionar(TIPO='AREA',LOCAL=r['LOCAL'],ID=legivel(p.get('id_original')),CAMADA=p['camada'],
                AREA=legivel(exibicao),
                TEC_UID=p['chave'],INFORMACAO=f"{fonte.get('modo')} | Campo exibição: {fonte.get('campo_nome')} | {trans}")
        grupos.append((principal,principal+1,len(linhas)+1))
    ws=planilha('DETALHAMENTO_LOCAIS',CAMPOS,linhas,'',estruturada=False,filtro=False)
    from openpyxl.styles import PatternFill,Font
    ws.sheet_properties.outlinePr.summaryBelow=False
    ws.sheet_properties.outlinePr.showOutlineSymbols=True
    ws.sheet_view.showOutlineSymbols=True
    ws.sheet_format.outlineLevelRow=1 if grupos else 0
    for principal,inicio,fim in grupos:
        for c in ws[principal]:
            c.fill=PatternFill('solid',fgColor='D9E2F3')
            c.font=Font(name='Calibri',size=11,bold=True,color='17365D')
        ws.row_dimensions[principal].collapsed=True
        if inicio<=fim:
            ws.row_dimensions.group(inicio,fim,outline_level=1,hidden=True)
    ws.column_dimensions['A'].width=16
    ws.column_dimensions['B'].width=12
    ws.column_dimensions['C'].width=24
    ws.column_dimensions['N'].width=65
    ws.column_dimensions['O'].hidden=True
    ws.print_title_rows='1:1'
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.orientation='landscape'
    ws.page_setup.paperSize=ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth=1
    ws.page_setup.fitToHeight=0
    ws.print_options.horizontalCentered=True
    ws.print_area=f'A1:N{ws.max_row}'
    return ws
