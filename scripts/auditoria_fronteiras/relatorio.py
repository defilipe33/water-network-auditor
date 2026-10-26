"""Relatório opcional para execução no Python do QGIS, sem dependência de Node.

openpyxl é importado somente ao exportar. Não instala pacotes nem sobrescreve
arquivos. O JSON acompanhante é a cópia completa e sem limite de célula Excel.
"""
from dataclasses import asdict
from io import BytesIO
from pathlib import Path
from math import ceil

from .apresentacao import CAMPOS_PRINCIPAIS, preparar_apresentacao, legivel
from .visualizacao import ESTILOS, resumir, texto_json
from .investigacao import (analisar_investigacao,validar_analise,
                           CAMPOS_LOCAIS,CAMPOS_PARES)


def exportar_xlsx(resultado, caminho, incluir_detalhes=False, analise=None, dossie=None):
    if resultado.cancelado:
        raise ValueError('Resultado parcial/cancelado não pode gerar relatório completo')
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.worksheet.table import Table, TableStyleInfo
        from openpyxl.worksheet.table import TableColumn
        from openpyxl.worksheet.filters import AutoFilter
        from openpyxl.utils import get_column_letter
    except ImportError as erro:
        raise RuntimeError('openpyxl não está disponível no Python do QGIS. '
                           'As camadas temporárias continuam disponíveis; nenhum pacote foi instalado.') from erro
    caminho = Path(caminho).resolve()
    if caminho.suffix.lower() != '.xlsx':
        raise ValueError('Escolha um nome com extensão .xlsx')
    tecnico = caminho.with_name(caminho.stem + '_tecnico.json')
    if not caminho.parent.is_dir():
        raise ValueError('A pasta de destino não existe')
    for p in (caminho, tecnico):
        if p.exists():
            raise FileExistsError(f'O arquivo já existe; escolha outro nome: {p.name}')
    itens = preparar_apresentacao(resultado)
    resumo = resumir(resultado)
    analise=analise or analisar_investigacao(resultado)
    validar_analise(resultado,analise)
    from .inspecao import construir_dossie,contagens
    dossie=dossie or construir_dossie(resultado,analise)
    if {e['uid'] for e in dossie['eventos']}!={e.id for e in resultado.ocorrencias}:
        raise ValueError('Dossiê não corresponde aos eventos deste relatório')
    counts=contagens(dossie)
    wb = Workbook()
    wb.remove(wb.active)

    def planilha(nome, cabecalhos, linhas, identificador,estruturada=True,filtro=True):
        cabecalhos = tuple(cabecalhos)
        if (not cabecalhos or any(not isinstance(c,str) or not c.strip() for c in cabecalhos)
                or len({c.casefold() for c in cabecalhos}) != len(cabecalhos)):
            raise ValueError('Cabeçalhos da tabela devem ser textos não vazios e únicos')
        ws = wb.create_sheet(nome)
        ws.sheet_view.showGridLines = False
        ws.append(list(cabecalhos))
        for valores in linhas:
            if len(valores) != len(cabecalhos):
                raise ValueError(f'Linha com quantidade de colunas incompatível na aba {nome}')
            if ws.max_row >= 1048576:
                raise ValueError('Quantidade de registros excede o limite de linhas do Excel')
            if any(isinstance(v,str) and len(v)>32767 for v in valores):
                raise ValueError(f'Texto excede o limite de célula Excel na aba {nome}; nenhum dado foi truncado')
            ws.append(list(valores))
        for linha in ws.iter_rows():
            for celula in linha:
                celula.font = Font(name='Calibri', size=11)
                celula.alignment = Alignment(vertical='top', wrap_text=True)
                if isinstance(celula.value, str):
                    # Preserva zeros à esquerda e impede tratar cadastros como fórmulas.
                    celula.data_type = 's'
                    celula.number_format = '@'
                elif isinstance(celula.value, (float,int)):
                    celula.number_format = '0.00' if cabecalhos[celula.column-1] in ('X','Y') else '#,##0'
                    if cabecalhos[celula.column-1]=='VALOR' and isinstance(celula.value,float):
                        celula.number_format='0.############'
        for celula in ws[1]:
            celula.fill = PatternFill('solid', fgColor='203864')
            celula.font = Font(name='Calibri',size=11,bold=True,color='FFFFFF')
            celula.alignment = Alignment(horizontal='center',vertical='center',wrap_text=True)
        ws.row_dimensions[1].height = 32
        ws.freeze_panes = 'A2'
        # O intervalo é o bloco de dados, não a dimensão da planilha (que pode
        # incluir notas fora da tabela). Um único proprietário do filtro:
        # tabela com dados OU worksheet sem tabela, nunca ambos sobrepostos.
        intervalo = f'A1:{get_column_letter(len(cabecalhos))}{ws.max_row}'
        if ws.max_row > 1 and estruturada:
            tabela = Table(displayName=identificador,ref=intervalo,
                tableColumns=[TableColumn(id=i,name=c) for i,c in enumerate(cabecalhos,1)],
                autoFilter=AutoFilter(ref=intervalo), totalsRowShown=False)
            tabela.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2',showRowStripes=True)
            ws.add_table(tabela)
        elif filtro:
            ws.auto_filter.ref = intervalo
        for i, nome_campo in enumerate(cabecalhos,1):
            largura = {'ID':8, 'CLASSIF':28, 'TIPO_LIM':42, 'AREA':38,
                       'REG_ANTES':38, 'REG_DEPOIS':38, 'REDE':42, 'ID_REDE':24,
                       'AMBIGUO':12, 'X':16, 'Y':16, 'SECAO':22, 'ITEM':38,
                       'QUANTIDADE':16, 'ID_AREA':22, 'EVT_UID':28,
                       'PARTE':10, 'JSON':80,'LOCAL':12,'N_EVENTOS':14,'N_TUBOS':14,
                       'TIPOS_LIM':40,'AREAS':38,'REGIOES':42,'SITUACAO':16,
                       'LEITURA':20,'REGIAO_A':38,'REGIAO_B':38,'EVENTOS':28,
                       'PARAMETRO':35,'VALOR':70,'CAMADA':42}.get(nome_campo,25)
            ws.column_dimensions[get_column_letter(i)].width = largura
        for linha in ws.iter_rows(min_row=2):
            linhas_texto = max((sum(max(1,ceil(len(p)/(ws.column_dimensions[c.column_letter].width*.9)))
                                   for p in str(c.value or '').split('\n')) for c in linha),default=1)
            ws.row_dimensions[linha[0].row].height = min(400,max(22,linhas_texto*15+6))
        return ws

    linhas_resumo = [('TOTAL','Ocorrências únicas',resumo['total'],'','',''),
        ('TOTAL','Eventos ambíguos',counts['eventos_ambiguos'],'','',''),
        ('MODO','Transição entre regiões',resumo['transicoes'],'','',''),
        ('MODO','Perímetro individual',resumo['perimetros'],'','','')]
    linhas_resumo.append(('INVESTIGACAO','Locais com ambiguidade',counts['locais_ambiguos'],'','',''))
    linhas_resumo += [('INVESTIGACAO',nome,analise.contadores[chave],'','','') for nome,chave in [
        ('Locais de investigação','locais'),('Tubulações distintas nos eventos','tubulacoes'),
        ('Pares inequívocos distintos','pares_inequivocos'),
        ('Ocorrências em pares inequívocos','ocorrencias_par_inequivoco'),
        ('Ocorrências com interpretação múltipla/indeterminada','interpretacoes_multiplas'),
        ('Travessias região/exterior inequívocas','saidas_externas')]]
    linhas_resumo += [('CLASSIFICACAO',c,resumo['classes'].get(c,0),'','','') for c in ESTILOS]
    linhas_resumo += [('CAMADA',c['camada'],c['total'],c['camada'],'','') for c in resumo['camadas']]
    linhas_resumo += [('AREA',a['area'],a['total'],a['camada'],a['area'],legivel(a['id_original']))
                      for a in resumo['areas']]
    ws = planilha('RESUMO',('SECAO','ITEM','QUANTIDADE','TIPO_LIM','AREA','ID_AREA'),
                   linhas_resumo,'ResumoAuditoria')
    notas = ['Critérios e parâmetros',
             'Uma ocorrência por ID técnico. Os modos incluem todas as classificações geométricas.',
             'Por área, cada ocorrência conta uma vez em cada polígono relacionado. A soma por áreas pode exceder o total.',
             'Ambígua inclui o indicador do núcleo, múltiplos candidatos e identificação de área indeterminada.',
             'Antes/depois é orientação geométrica, sem significado de sentido hidráulico.',
             f'CRS: {resultado.parametros.crs}',
             f'Tolerância geométrica (m): {resultado.parametros.tolerancia_geometrica}',
             f'Distância de continuidade (m): {resultado.parametros.distancia_continuidade}',
             f'QGIS: {resultado.ambiente.get("qgis", "")}',
             f'Avisos do processamento: {len(resultado.problemas)}',
             'Rastreabilidade integral: ' + tecnico.name]
    for i, nota in enumerate(notas,1):
        ws.cell(i,8,nota).alignment = Alignment(wrap_text=True,vertical='top')
        ws.row_dimensions[i].height = max(ws.row_dimensions[i].height or 15, 45)
    ws.column_dimensions['H'].width = 70
    for modo,nome,tabela in [('TRANSICAO','TRANSICOES_SETORES','TransicoesAuditoria'),
                              ('PERIMETRO','TRAVESSIAS_PERIMETROS','PerimetrosAuditoria')]:
        planilha(nome,CAMPOS_PRINCIPAIS,
            [list(i['registro'].values()) for i in itens if i['modo']==modo],tabela)
    planilha('PONTOS_INVESTIGACAO',CAMPOS_LOCAIS,
        [[l['registro'][c] for c in CAMPOS_LOCAIS] for l in dossie['locais']],'LocaisAuditoria')
    planilha('PARES_SETORES',CAMPOS_PARES,
        [[p[c] for c in CAMPOS_PARES] for p in analise.pares],'ParesAuditoria')
    planilha('SAIDAS_EXTERNAS',CAMPOS_PRINCIPAIS,
        [[s['registro'][c] for c in CAMPOS_PRINCIPAIS] for s in analise.saidas],'ExternasAuditoria')
    parametros=[('EXECUCAO','',n,str(v)) for n,v in [
        ('ID técnico da execução',dossie['execucao_uid']),
        ('Data da execução (UTC)',analise.metadados['data_execucao_utc'] if
            analise.metadados['data_execucao_utc']!='Não registrada' else
            dossie.get('controle_versao',{}).get('data_processamento_utc','Não registrada')),
        ('Data do pós-processamento (UTC)',analise.metadados['data_posprocessamento_utc']),
        ('Versão dos locais',analise.metadados['versao']),('CRS processamento',resultado.parametros.crs),
        ('QGIS',resultado.ambiente.get('qgis',''))]]
    parametros += [('NUCLEO','',n,v) for n,v in asdict(resultado.parametros).items()]
    from .versoes import metadados_exportacao
    versao_exportada=metadados_exportacao(dossie)
    parametros += [('VERSAO','',n,v if v is not None else '')
        for n,v in versao_exportada['versao'].items()]
    parametros += [('EXPORTACAO','',n,v if v is not None else '')
        for n,v in versao_exportada.items() if n not in ('parametros','versao')]
    parametros += [('AGRUPAMENTO','',n,v) for n,v in asdict(analise.opcoes).items()]
    for origem,campos in dossie['campos_interesse'].items():
        parametros += [('INSPECAO',origem,n,v or '(Não escolhido)') for n,v in campos.items()]
    for fonte in resultado.fontes:
        categoria='LIMITE' if fonte.get('modo') else 'REDE'
        for n,valor in fonte.items():
            if n not in ('origem','camada'):
                parametros.append((categoria,fonte['camada'],
                    'Campo de exibição' if n=='campo_nome' else n,valor if valor is not None else '(Não escolhido)'))
        contagem=next((v for k,v in resultado.contadores.items()
                       if k.endswith(':'+fonte['origem']) and isinstance(v,dict)),{})
        parametros += [(categoria,fonte['camada'],'Contagem: '+n,v if v is not None else '(Não disponível)')
                        for n,v in contagem.items()]
    parametros += [('AVISO',p.get('camada',''),str(n),p.get('mensagem',str(p)))
                   for n,p in enumerate(resultado.problemas,1)]
    planilha('PARAMETROS',('SECAO','CAMADA','PARAMETRO','VALOR'),parametros,'ParametrosAuditoria')
    from .relatorio_detalhado import criar_detalhamento
    criar_detalhamento(wb,dossie,planilha)
    if incluir_detalhes:
        linhas = []
        for item in itens:
            dados = texto_json(item['tecnico'])
            for parte,inicio in enumerate(range(0,len(dados),30000),1):
                linhas.append([item['registro']['ID'],item['uid'],parte,dados[inicio:inicio+30000]])
        planilha('DETALHES_TECNICOS',('ID','EVT_UID','PARTE','JSON'),linhas,'DetalhesAuditoria')
    buffer = BytesIO()
    wb.save(buffer)
    dados_tecnicos = texto_json(dict(versao=1, parametros=asdict(resultado.parametros),
        ambiente=resultado.ambiente, fontes=resultado.fontes, contadores=resultado.contadores,
        problemas=resultado.problemas, cancelado=resultado.cancelado,
        ocorrencias=[dict(ID=i['registro']['ID'],uid=i['uid'],registro=i['registro'],
                          tecnico=i['tecnico']) for i in itens],investigacao=asdict(analise),
        dossie=dossie,exportacao=versao_exportada)).encode('utf-8')
    criados = []
    try:
        for p, dados in ((tecnico,dados_tecnicos),(caminho,buffer.getvalue())):
            with p.open('xb') as arquivo:
                criados.append(p)
                arquivo.write(dados)
    except Exception:
        for p in criados:
            p.unlink()  # Apenas arquivos novos criados nesta chamada.
        raise
    return dict(xlsx=str(caminho), tecnico=str(tecnico))
