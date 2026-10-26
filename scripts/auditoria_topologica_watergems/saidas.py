"""Resultados independentes do banco original, sem sobrescrita implícita."""
from collections import Counter
from datetime import datetime, timezone
import json
from pathlib import Path
from uuid import uuid4

DICIONARIO_CAMPOS={
    'WG_ID':'DomainElementID original, preservado como texto',
    'WG_UID':'Identidade técnica: projeto + DomainElementID',
    'ROTULO':'Label original, não necessariamente único',
    'TIPO':'Tipo/família do elemento WaterGEMS',
    'ATIVO':'Active Topology resolvida no cenário; desconhecido não equivale a ativo',
    'EXCLUIDO':'Exclusão confrontada no cadastro; feições exportadas são vivas',
    'CENARIO':'ScenarioID selecionado',
    'HIPOTESE':'Perspectiva estrutural ou hipótese; não representa vazão/estado real',
    'COMPONENTE':'Número de apresentação do componente nesta execução',
    'COMP_UID':'Identificador estável da composição técnica/cenário/perspectiva',
    'NO_INICIO':'Extremidade Start resolvida, em DomainElementID',
    'NO_FIM':'Extremidade Stop resolvida, em DomainElementID',
    'LOCAL_UID':'Identidade original do local de investigação do dossiê',
    'EVT_UID':'Identidade original do evento de fronteira',
    'GIS_UID':'Identidade técnica GIS; não é DomainElementID',
    'CHAVE':'Valor cadastral do campo explicitamente mapeado (ID no ensaio)',
    'SITUACAO':'Qualidade da correspondência cadastral/geométrica, sem confirmação hidráulica',
    'ERRO_M':'Maior divergência XY entre vértices, em metros',
    'NATUREZA':'Incidência explícita ou hipótese geométrica, sem mistura',
    'PROPORCAO':'Fração das tubulações aceitas na perspectiva pertencente ao componente'}


def serializavel(v):
    if isinstance(v,bytes): return {'wkb_hex':v.hex(), 'bytes':len(v)}
    raise TypeError(type(v).__name__)


def tabelas(modelo, resultado):
    resumo=resultado['resumo']
    inventario=[]
    for i,e in modelo['elementos'].items():
        inventario.append((str(i),e['label'],e['tipo'],e['familia'],e['excluido'],e['ativo'],e['uid']))
    abas={
        'RESUMO':(('INDICADOR','VALOR'),[(k,v) for k,v in resumo.items() if not isinstance(v,dict)]+[
            ('Advertência',resultado['aviso']),('Validação WaterGEMS desktop','PENDENTE'),
            ('Fonte',modelo['banco'])]),
        'TIPOS':(('TIPO','QUANTIDADE'),list(resumo['inventario'].items())),
        'ELEMENTOS':(('ID','ROTULO','TIPO','FAMILIA','EXCLUIDO','ATIVO','UID'),inventario),
        'CONEXOES':(('TUBO','NO_INICIO','NO_FIM'),[(str(c['id']),str(c['inicio']),str(c['fim'])) for c in resultado['conexoes']]),
        'COMPONENTES':(('COMPONENTE','UID','NOS','TUBOS','DISPOSITIVOS','FONTES','REF_PEND','PROPORCAO','TAMANHO','EXEMPLOS'),[
            (c['id'],c.get('uid',''),c.get('n_nos',len(c['nos'])),len(c['tubulacoes']),c.get('n_dispositivos',0),
             '; '.join(map(str,c.get('fontes',[]))),c.get('referencias_nao_resolvidas',0),c.get('proporcao_tubos',0),
             c.get('classe_tamanho',''),'; '.join(str(e['id'])+' — '+str(e['label']) for e in c.get('representantes',[]))) for c in resultado['componentes']]),
        'DIAGNOSTICOS':(('CODIGO','ID','RELACIONADO','DESCRICAO'),[(d['codigo'],str(d['id']),'' if d.get('relacionado') is None else str(d['relacionado']),d.get('descricao','')) for d in resultado['diagnosticos']]),
        'PARAMETROS':(('PARAMETRO','VALOR'),[('Cenario',str(modelo['cenario'])),('Hipotese',resumo['modo'])]+[
            ('Alternativa '+k,'; '.join(map(str,v))) for k,v in modelo['alternativas'].items()]+[
            ('Aviso',a) for a in modelo['avisos']]+[('Referencia '+k,str(v)) for k,v in modelo.get('referencia_codigo',{}).items()])}
    if 'auditoria_jusantes' in resultado:
        abas['JUSANTES']=(('ID','LABEL','TIPO','JUSANTE','EXISTE','LINK_EXCL','LINK_ATIVO','DISP_ATIVO','CENARIO','ALT_DISP','ALT_LINK','CLASSIFICACAO','INCIDENTES','CAND_GEOM'),[
            (str(r['id']),r['label'],r['tipo'],None if r['jusante'] is None else str(r['jusante']),r['existe'],r['link_excluido'],r['link_ativo'],r['dispositivo_ativo'],str(r['cenario']),r['alternativa'],r['alternativa_link'],r['classificacao'],
             '; '.join(map(str,r['incidentes'])),'; '.join(str(c['link']) for c in r.get('candidatos_geometricos',[]))) for r in resultado['auditoria_jusantes']])
    if 'associacoes' in resultado:
        abas['ASSOCIACOES']=(('GIS_UID','ORIGEM','CHAVE','SITUACAO','CANDIDATOS','ERRO_XY_M'),[
            (r['gis_uid'],r['origem'],r['chave_cadastral'],r['situacao'],'; '.join(map(str,r['candidatos'])),r['erro_xy_m']) for r in resultado['associacoes']])
    if 'validacao_cenarios' in resultado:
        abas['CENARIOS']=(('CENARIO','TIPO_ALT','TABELA','ELEMENTO','ALT_LEITOR','ALT_COMPOSER','IGUAL','FONTE_CENARIO'),[
            (r['cenario'],r['tipo'],r['tabela'],str(r['elemento']),r['alternativa_leitor'],r['alternativa_composer'],r['igual'],r['cenario_fonte']) for r in resultado['validacao_cenarios']])
    if 'comparacao_modos' in resultado:
        modos=resultado['comparacao_modos']
        abas['COMPARACAO']=(('MODO','COMPONENTES','CONEXOES','INATIVOS','BOMBAS_INDET','LIMITACAO'),[
            (s['modo'],s['componentes'],s['conexoes'],s['diagnosticos'].get('INATIVO',0),s['diagnosticos'].get('BOMBA_ESTADO_INDETERMINADO',0),s['restricoes'])
            for s in (modos['estrutural'],modos['todas_abertas'])]+[('ESTADOS_DO_CENARIO',None,None,None,None,modos['estados_cenario'])])
    return abas


def exportar_json(modelo, resultado, destino, completo=True):
    p=Path(destino)
    dados=dict(execucao_uid=str(uuid4()),data=datetime.now(timezone.utc).isoformat(),
        modelo={k:v for k,v in modelo.items() if completo or k!='elementos'},
        resultado={k:v for k,v in resultado.items() if k!='grafo'})
    with p.open('x',encoding='utf-8') as f:
        json.dump(dados,f,ensure_ascii=False,default=serializavel)
    return p


def exportar_xlsx(modelo, resultado, destino):
    return exportar_tabelas_xlsx(tabelas(modelo,resultado),destino)


def exportar_tabelas_xlsx(abas,destino,larguras=None):
    """Mesmo exportador OOXML para relatórios topológicos e integrados."""
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font,PatternFill,Alignment
    from openpyxl.worksheet.table import Table,TableColumn,TableStyleInfo
    from openpyxl.worksheet.filters import AutoFilter
    from openpyxl.utils import get_column_letter
    p=Path(destino)
    if p.exists(): raise FileExistsError(str(p))
    wb=Workbook(write_only=True)
    for n,(nome,(cabecalhos,linhas)) in enumerate(abas.items(),1):
        w=wb.create_sheet(nome);w.freeze_panes='A2';w.sheet_view.showGridLines=False
        w.row_dimensions[1].height=30
        for k,h in enumerate(cabecalhos,1):
            w.column_dimensions[get_column_letter(k)].width=(larguras or {}).get(h,60 if h in ('DESCRICAO','VALOR','UID') else 26)
        titulos=[]
        for h in cabecalhos:
            c=WriteOnlyCell(w,value=h);c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='203864')
            c.alignment=Alignment(wrap_text=True);titulos.append(c)
        w.append(titulos);n_linhas=1
        for linha in linhas:
            if len(linha)!=len(cabecalhos): raise ValueError('Linha inconsistente')
            if n_linhas>=1048576: raise ValueError('Limite Excel excedido; use JSON completo')
            if any(isinstance(v,str) and len(v)>32767 for v in linha): raise ValueError('Célula extensa; nenhum texto truncado')
            celulas=[]
            if larguras:
                from math import ceil
                altura=max((ceil(len(str(v or ''))/max(10,w.column_dimensions[get_column_letter(k)].width*.85)) for k,v in enumerate(linha,1)),default=1)
                w.row_dimensions[n_linhas+1].height=min(120,max(24,altura*15))
            for v in linha:
                c=WriteOnlyCell(w,value=v)
                if isinstance(v,str):c.data_type='s'
                if larguras:c.alignment=Alignment(vertical='top',wrap_text=True)
                celulas.append(c)
            w.append(celulas);n_linhas+=1
        ref=f'A1:{get_column_letter(len(cabecalhos))}{n_linhas}'
        if linhas:
            t=Table(displayName=f'Topologia{n}',ref=ref,
                tableColumns=[TableColumn(id=k,name=h) for k,h in enumerate(cabecalhos,1)],
                autoFilter=AutoFilter(ref=ref),totalsRowShown=False)
            t.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True);w.add_table(t)
        else: w.auto_filter.ref=ref
    with p.open('xb') as f: wb.save(f)
    return p


def camadas(modelo, resultado, geoms=None, crs=None):
    """CRS ausente: geometria bruta em referência desconhecida, apenas GPKG."""
    from qgis.core import QgsVectorLayer,QgsField,QgsFeature,QgsGeometry,QgsWkbTypes
    from qgis.PyQt.QtCore import QVariant
    saida=[]
    for familia,tipo,nome in [('LINK','LineString','WG_TUBULACOES'),('NO','Point','WG_NOS'),('DISPOSITIVO','Point','WG_DISPOSITIVOS')]:
        c=QgsVectorLayer(tipo+(f'?crs={crs}' if crs else ''),nome,'memory')
        if not crs:
            from qgis.core import QgsCoordinateReferenceSystem
            c.setCrs(QgsCoordinateReferenceSystem())
        c.dataProvider().addAttributes([QgsField(n,QVariant.String,len=l) for n,l in [('WG_ID',30),('WG_UID',80),('ROTULO',254),('TIPO',60),('ATIVO',12),('EXCLUIDO',12),('CENARIO',30),('HIPOTESE',40)]])
        c.updateFields();fs=[]
        for i,e in modelo['elementos'].items():
            if e['familia']!=familia or e['excluido'] is not False: continue
            g=(geoms or {}).get(i)
            if g is None and e['geometria']:
                g=QgsGeometry();g.fromWkb(bytes(e['geometria']))
            f=QgsFeature(c.fields());f.setAttributes([str(i),e['uid'],e['label'],e['tipo'],str(e['ativo']),str(e['excluido']),str(modelo['cenario']),resultado['resumo']['modo']])
            if g and not g.isEmpty(): f.setGeometry(g)
            fs.append(f)
        c.dataProvider().addFeatures(fs);c.updateExtents();c.setCustomProperty('auditoria_topologica_watergems/resultado',True)
        saida.append(c)
    # Tabelas sem geometria acompanham as saídas geográficas; uma linha por
    # vínculo. Não comprimir relacionamentos em textos paralelos truncáveis.
    for nome,campos,linhas in [
        ('WG_DIAGNOSTICOS',['CODIGO','WG_ID','REL_ID','DESCRICAO'],
         [(d['codigo'],str(d['id']),None if d.get('relacionado') is None else str(d['relacionado']),d.get('descricao','')) for d in resultado['diagnosticos']]),
        ('WG_CONEXOES',['WG_ID','NO_INICIO','NO_FIM'],
         [(str(x['id']),str(x['inicio']),str(x['fim'])) for x in resultado['conexoes']]),
        ('WG_COMPONENTES',['COMPONENTE','WG_ID','FAMILIA','COMP_UID'],
         [(str(g['id']),str(i),tipo,g.get('uid','')) for g in resultado['componentes'] for tipo,ids in [('NO',g['nos']),('LINK',g['tubulacoes'])] for i in ids])]:
        c=QgsVectorLayer('None',nome,'memory')
        c.dataProvider().addAttributes([QgsField(n,QVariant.String,len=254) for n in campos]);c.updateFields()
        fs=[]
        for valores in linhas:
            f=QgsFeature(c.fields());f.setAttributes(list(valores));fs.append(f)
        c.dataProvider().addFeatures(fs);saida.append(c)
    extras=[('WG_COMP_INFO',['COMPONENTE','COMP_UID','N_NOS','N_TUBOS','N_DISP','N_FONTES','REF_PEND','PROPORCAO','TAMANHO'],[
        [str(c['id']),c.get('uid',''),str(c.get('n_nos',0)),str(len(c['tubulacoes'])),str(c.get('n_dispositivos',0)),str(len(c.get('fontes',[]))),
         str(c.get('referencias_nao_resolvidas',0)),repr(c.get('proporcao_tubos',0)),c.get('classe_tamanho','')] for c in resultado['componentes']])]
    if 'auditoria_jusantes' in resultado:
        extras.append(('WG_JUSANTES',['WG_ID','ROTULO','TIPO','JUS_ID','CLASSIF','EXISTE','EXCLUIDO','ATIVO','CENARIO','ALT_ID'],[
            [str(r['id']),r['label'],r['tipo'],None if r['jusante'] is None else str(r['jusante']),r['classificacao'],str(r['existe']),str(r['link_excluido']),str(r['link_ativo']),str(r['cenario']),str(r['alternativa'])]
            for r in resultado['auditoria_jusantes']]))
        extras.append(('WG_JUS_REL',['DISP_ID','TUBO_ID','NATUREZA'],[
            [str(r['id']),str(i),'INCIDENCIA_EXPLICITA'] for r in resultado['auditoria_jusantes'] for i in r['incidentes']]+[
            [str(r['id']),str(c['link']),'HIPOTESE_GEOMETRICA_BRUTA'] for r in resultado['auditoria_jusantes'] for c in r.get('candidatos_geometricos',[])]))
    if 'associacoes' in resultado:
        extras.append(('WG_ASSOC',['GIS_UID','ORIGEM','CHAVE','WG_ID','SITUACAO','ERRO_M'],[
            [r['gis_uid'],r['origem'],r['chave_cadastral'],None if i is None else str(i),r['situacao'],None if r['erro_xy_m'] is None else repr(r['erro_xy_m'])]
            for r in resultado['associacoes'] for i in (r['candidatos'] or [None])]))
    if 'vinculos_fronteiras' in resultado:
        extras.append(('WG_FRONTEI',['LOCAL_UID','EVT_UID','GIS_UID','WG_ID','SITUACAO'],[
            [r['local_uid'],r['evento_uid'],r['gis_uid'],None if i is None else str(i),r['situacao']]
            for r in resultado['vinculos_fronteiras'] for i in (r['candidatos'] or [None])]))
    for nome,campos,linhas in extras:
        c=QgsVectorLayer('None',nome,'memory');c.dataProvider().addAttributes([QgsField(n,QVariant.String,len=254) for n in campos]);c.updateFields()
        fs=[]
        for valores in linhas:
            f=QgsFeature(c.fields());f.setAttributes(valores);fs.append(f)
        c.dataProvider().addFeatures(fs);saida.append(c)
    return saida


def exportar_geografico(camadas, pasta, formato='GPKG',campo_chave='WG_UID'):
    from qgis.core import QgsVectorFileWriter,QgsProject,QgsVectorLayer
    pasta=Path(pasta);pasta.mkdir(parents=True,exist_ok=False)
    if formato not in ('GPKG','SHP'): raise ValueError('Formato desconhecido')
    resultados=[]
    for k,c in enumerate(camadas):
        if not c.isSpatial() and formato=='SHP':
            import csv
            arq=pasta/(c.name()+'.csv')
            with arq.open('x',encoding='utf-8-sig',newline='') as f:
                w=csv.writer(f,delimiter=';');w.writerow(c.fields().names())
                for e in c.getFeatures():w.writerow(e.attributes())
            resultados.append(dict(arquivo=str(arq),camada=c.name(),registros=c.featureCount(),reimportacao='CSV auxiliar, sem truncamento DBF'))
            continue
        if formato=='SHP' and not c.crs().isValid():
            raise ValueError('SHP exige CRS confirmado; use GPKG bruto sem referência ou informe CRS/unidade')
        for f in c.fields():
            if formato=='SHP' and (len(f.name())>10 or f.length()>254): raise ValueError('Campo DBF incompatível')
        if formato=='SHP':
            for f in c.getFeatures():
                for a,campo in zip(f.attributes(),c.fields()):
                    if a is not None and len(str(a).encode('utf-8'))>254:
                        raise ValueError('Texto excede 254 bytes DBF; utilize GPKG, sem truncamento silencioso')
        arq=pasta/('TOPOLOGIA.gpkg' if formato=='GPKG' else c.name()+'.shp')
        o=QgsVectorFileWriter.SaveVectorOptions();o.driverName='GPKG' if formato=='GPKG' else 'ESRI Shapefile'
        o.fileEncoding='UTF-8';o.layerName=c.name()
        if formato=='GPKG' and k: o.actionOnExistingFile=QgsVectorFileWriter.CreateOrOverwriteLayer
        r=QgsVectorFileWriter.writeAsVectorFormatV3(c,str(arq),QgsProject.instance().transformContext(),o)
        if r[0]!=QgsVectorFileWriter.NoError: raise RuntimeError(str(r))
        nova=QgsVectorLayer(str(arq)+(f'|layername={c.name()}' if formato=='GPKG' else ''),c.name(),'ogr')
        if not nova.isValid() or nova.featureCount()!=c.featureCount(): raise ValueError('Falha de reimportação/contagem')
        def vertices(g):
            return [(v.x(),v.y())+((v.z(),) if g.constGet() and g.constGet().is3D() else ())
                    +((v.m(),) if g.constGet() and g.constGet().isMeasure() else ())
                    for v in g.vertices()]
        if c.isSpatial():
            esperado={f[campo_chave]:vertices(f.geometry()) for f in c.getFeatures()}
            for f in nova.getFeatures():
                if f[campo_chave] not in esperado or vertices(f.geometry())!=esperado[f[campo_chave]]:
                    raise ValueError('Identidade/geometria alterada na exportação')
        else:
            if [f.attributes() for f in nova.getFeatures()] and nova.fields().names()[0]=='fid':
                preservados=[f.attributes()[1:] for f in nova.getFeatures()]
            else:preservados=[f.attributes() for f in nova.getFeatures()]
            if preservados!=[f.attributes() for f in c.getFeatures()]:raise ValueError('Relacionamentos alterados')
        if c.crs().isValid() and nova.crs()!=c.crs(): raise ValueError('CRS alterado')
        resultados.append(dict(arquivo=str(arq),camada=c.name(),registros=nova.featureCount(),reimportacao='OK'))
        del nova
    (pasta/'VALIDACAO.json').write_text(json.dumps(resultados,ensure_ascii=False,indent=2),encoding='utf-8')
    (pasta/'DICIONARIO_CAMPOS.json').write_text(json.dumps(DICIONARIO_CAMPOS,ensure_ascii=False,indent=2),encoding='utf-8')
    return resultados
